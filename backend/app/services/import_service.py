from __future__ import annotations

import json
import hashlib
import re
import unicodedata
from datetime import datetime, timezone
from difflib import SequenceMatcher
from io import BytesIO
from pathlib import PurePosixPath
from time import perf_counter

import fitz
import httpx
import pytesseract
from fastapi import HTTPException, UploadFile, status
from openpyxl import load_workbook
from PIL import Image
from sqlalchemy import text
from sqlalchemy.orm.attributes import flag_modified
from sqlalchemy.orm import Session

from app.models.entities import (
    AnnualGoal,
    AcademicPeriod,
    Career,
    ExternalResearcher,
    ImportJob,
    ImportNormalizationAudit,
    ImportReviewItem,
    ImportedOcrTrace,
    ImportedProgressReport,
    ImportedProjectParticipant,
    ImportedResearchRecord,
    PersonRole,
    ProjectTeacher,
    ResearchEntity,
    ResearchProject,
    ScientificProduction,
    ScientificProductionAuthor,
    Teacher,
)
from app.core.config import settings
from app.models.enums import GoalMetric, ProductionType, ProjectTeacherRole, ProjectType
from app.schemas.imports import (
    DropboxProgressPdfBatchRequest,
    DropboxProgressPdfItem,
    ImportResult,
    OcrTraceReviewUpdate,
    ProgressJsonImportRequest,
)
from app.services.import_batching import dropbox_document_key, dropbox_fingerprint, promote_document_version
from app.services.import_reconciliation import clean_product_title, product_title_quality
from app.services.institution_normalizer import normalize_external_institution_display
from app.services.person_name_matching import heuristic_name_merge_allowed
from app.services.pdf_parser import classify_document, parse_progress_report
from app.services.pdf_parser.pdf_text_extractor import extract_pdf_text
from app.services.pdf_parser.pdf_text_extractor import PdfTextExtraction


def _normalize_header(value: object) -> str:
    text = str(value or "").strip().upper()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    return " ".join(text.split())


def _utcnow() -> datetime:
    return datetime.now(timezone.utc)


def _elapsed_ms(start: datetime | None, end: datetime | None = None) -> int | None:
    if not start:
        return None
    start_utc = start.replace(tzinfo=timezone.utc) if start.tzinfo is None else start.astimezone(timezone.utc)
    end_value = end or _utcnow()
    end_utc = end_value.replace(tzinfo=timezone.utc) if end_value.tzinfo is None else end_value.astimezone(timezone.utc)
    return max(0, int((end_utc - start_utc).total_seconds() * 1000))


def _normalize_text(value: object) -> str:
    return str(value or "").strip()


def _normalize_key(value: object) -> str:
    text = _normalize_text(value).upper()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    return " ".join(text.split())


def _person_match_tokens(value: object) -> set[str]:
    particles = {"DE", "DEL", "LA", "LAS", "LOS", "Y"}
    return {
        token
        for token in _normalize_key(value).split()
        if token not in particles and len(token) > 1
    }


def _is_author_context(source_section: str | None, source_field: str | None, row_or_block_id: str | None) -> bool:
    section = str(source_section or "").strip().lower()
    field = str(source_field or "").strip().lower()
    return section in {"intercambios", "produccion_cientifica"} and field in {"author", "authors", "autor"} and bool(row_or_block_id)


def _person_token_match_score(left_tokens: set[str], right_tokens: set[str]) -> float:
    left = [token for token in left_tokens if len(token) > 1]
    right = [token for token in right_tokens if len(token) > 1]
    if len(left) < 2 or len(right) < 2:
        return 0.0
    shorter, longer = (left, right) if len(left) <= len(right) else (right, left)
    matched: set[str] = set()
    matches = 0
    for token in shorter:
        candidate = _best_person_token_match(token, longer, matched)
        if candidate:
            matched.add(candidate)
            matches += 1
    return matches / len(shorter)


def _best_person_token_match(token: str, candidates: list[str], used: set[str]) -> str | None:
    if token in candidates and token not in used:
        return token
    best_candidate = None
    best_score = 0.0
    for candidate in candidates:
        if candidate in used:
            continue
        score = SequenceMatcher(None, token, candidate).ratio()
        if score > best_score:
            best_candidate = candidate
            best_score = score
    return best_candidate if best_score >= 0.84 else None


def _optional_int(value: object) -> int | None:
    if value in (None, ""):
        return None
    try:
        return int(float(str(value)))
    except (TypeError, ValueError):
        return None


def _normalize_int(value: object, default: int = 0) -> int:
    if value in (None, ""):
        return default
    return int(float(value))


def _sanitize_db_text(value: str | None) -> str | None:
    if value is None:
        return None

    # PostgreSQL text columns reject NUL bytes. We also drop other control
    # characters except common whitespace to keep OCR content readable.
    cleaned = "".join(
        char for char in value if char in ("\n", "\r", "\t") or ord(char) >= 32
    )
    cleaned = cleaned.replace("\x00", "")
    return cleaned.strip() or None


def _roman_cycle_to_int(value: str | None) -> int:
    normalized = _normalize_key(value)
    if normalized in {"2", "II", "I I"}:
        return 2
    return 1


def _extract_number_before(text: str, words: list[str]) -> int:
    normalized = _normalize_key(text)
    for word in words:
        pattern = rf"(\d+)\s+{word}"
        match = re.search(pattern, normalized)
        if match:
            return int(match.group(1))
    return 0


GENERATED_PROGRESS_NOTE = "Investigacion reportada desde PDF."
LEGACY_GENERATED_PROGRESS_NOTE = "Registro preliminar generado desde texto OCR de PDF."
PARSER_VERSION = "section_context_v2"
NON_PERSON_NAME_WORDS = {
    "ADMINISTRACION",
    "ADMINISTRATIVAS",
    "AUTOR",
    "AUDITORIA",
    "CARRERA",
    "CIENCIAS",
    "COMERCIO",
    "CONTABILIDAD",
    "CORRECTAMENTE",
    "DOCENTE",
    "EDUCACION",
    "ELECTRONICAMENTE",
    "EMPRESAS",
    "ENVIAR",
    "EXTERIOR",
    "FACULTAD",
    "FILOSOFIA",
    "FIRMADO",
    "FORMULARIO",
    "GERENCIAL",
    "GESTION",
    "GRUPO",
    "INFORMACION",
    "INVESTIGADORES",
    "INVESTIGACION",
    "ALMERIA",
    "MURCIA",
    "UNIVERSIDAD",
    "INSTITUCION",
    "INSTITUCIONES",
    "LETRAS",
    "LICENCIATURA",
    "NOMBRE",
    "NACION",
    "DETECTADO",
    "DIRECTOR",
    "AGREGAR",
    "NUEVO",
    "NO",
    "POR",
    "PRESIONAR",
    "PROYECTO",
    "REGISTRO",
    "RESPONSABLE",
}


def _title_case_name(value: str) -> str:
    return " ".join(part.capitalize() for part in value.split())


def _person_key(value: str | None) -> str:
    key = _normalize_key(value)
    particles = {"DE", "DEL", "LA", "LAS", "LOS", "Y"}
    return " ".join(token for token in key.split() if token not in particles)


def _clean_person_name(value: str | None) -> str | None:
    cleaned, _ = _clean_ocr_person_text(value)
    if not _looks_like_person_name(cleaned):
        return None
    return _title_case_name(_normalize_key(cleaned or ""))


def _clean_ocr_person_text(value: str | None) -> tuple[str, str | None]:
    text = re.sub(r"\s+", " ", str(value or "").replace(".", " ").replace(":", " ")).strip(" -")
    tokens = _normalize_key(text).split()
    if len(tokens) >= 2 and len(tokens[0]) <= 2 and tokens[0] not in {"DE", "LA", "EL"}:
        return " ".join(tokens[1:]), "Nombre con prefijo OCR corto removido; posible nombre truncado."
    return text, None


def _join_reasons(*reasons: str | None) -> str:
    return " ".join(reason.strip() for reason in reasons if reason and reason.strip())


def _invalid_author_fragment(value: str | None) -> bool:
    tokens = _normalize_key(value).split()
    if not tokens:
        return True
    if any(token in NON_PERSON_NAME_WORDS for token in tokens):
        return True
    if any(len(token) <= 1 for token in tokens):
        return True
    return False


def _clean_career_name(value: str | None) -> str | None:
    key = _normalize_key(value)
    if not key:
        return None
    if "CONTABILIDAD" in key and "AUDITORIA" in key:
        return "Licenciatura en Contabilidad y Auditoria"
    if "ADMINISTRACION" in key and "EMPRESAS" in key:
        return "Licenciatura en Administracion de Empresas"
    if "COMERCIO" in key and "EXTERIOR" in key:
        return "Licenciatura en Comercio Exterior"
    if "MERCADOTECNIA" in key:
        return "Licenciatura en Mercadotecnia"
    if "FINANZAS" in key:
        return "Licenciatura en Finanzas"
    if "GESTION" in key and "INFORMACION" in key:
        return "Licenciatura en Gestion de la Informacion Gerencial"
    if "TURISMO" in key:
        return "Licenciatura en Turismo"
    if "NEGOCIOS" in key and "INTERNACIONALES" in key:
        return "Licenciatura en Negocios Internacionales"
    if "PEDAGOGIA" in key:
        return "Licenciatura en Pedagogia de las Ciencias Experimentales de las Matematicas"
    return None


def _clean_external_institution(value: str | None) -> str | None:
    text = " ".join(str(value or "").replace(":", " ").split()).strip(" -")
    institution = normalize_external_institution_display(value)
    if not institution:
        return None
    key = _normalize_key(institution)
    if _looks_like_person_name(text) or re.search(r"HTTP|WWW|JOTFORM|\d+\s+DE\s+\d+", key):
        return None
    return institution


def _external_participant_type(value: str | None, default: str | None = None) -> str:
    key = _normalize_key(value)
    if "ESTUDIANTE" in key or "ALUMNO" in key:
        return "estudiante"
    if "GRADUADO" in key or "EGRESADO" in key:
        return "graduado"
    if default:
        return default
    if any(marker in key for marker in ("UNIVERSIDAD", "INSTITUTO", "ESCUELA", "CENTRO", "POLITECNICA")):
        return "investigador_externo"
    return "no_clasificado"


def _is_pdf_noise_line(value: str | None) -> bool:
    key = _normalize_key(value)
    return (
        not key
        or "FORMULARIO DE INFORME" in key
        or "JOTFORM" in key
        or key.startswith("HTTPS")
        or bool(re.fullmatch(r"\d+\s+DE\s+\d+", key))
        or bool(re.fullmatch(r"\d{1,2}/\d{1,2}/\d{4}.*", key))
    )


def _looks_like_person_name(value: str | None) -> bool:
    if not value:
        return False

    cleaned = re.sub(r"\s+", " ", value.replace(".", " ").replace(":", " ")).strip(" -")
    key = _normalize_key(cleaned)
    tokens = key.split()
    if len(tokens) < 3 or len(tokens) > 6:
        return False
    if re.search(r"\d|/|\\|_|\.PDF|SIGNED|INFSEM|GI\d", key):
        return False
    if any(word in NON_PERSON_NAME_WORDS for word in tokens):
        return False
    if any(len(token) <= 1 for token in tokens):
        return False
    function_words = {"A", "AL", "COMO", "CON", "DE", "DEL", "EL", "EN", "LA", "LAS", "LOS", "PARA", "POR", "UN", "UNA", "Y"}
    meaningful_tokens = [token for token in tokens if token not in function_words]
    if len(meaningful_tokens) < 2 or any(len(token) < 3 for token in meaningful_tokens):
        return False
    if len(tokens) >= 4 and sum(1 for token in tokens if token in function_words) / len(tokens) >= 0.4:
        return False
    return all(re.fullmatch(r"[A-ZÑ]+", token) for token in tokens)


def _clean_topic(value: str | None) -> str | None:
    if not value:
        return None
    text = " ".join(value.replace(":", " ").split()).strip(" -")
    return text[:220] if len(text) >= 6 else None


def _collapse_repeated_words(words: list[str]) -> list[str]:
    if len(words) % 2 == 0:
        middle = len(words) // 2
        if [_normalize_key(item) for item in words[:middle]] == [_normalize_key(item) for item in words[middle:]]:
            return words[:middle]
    return words


def _is_section_heading(line: str) -> bool:
    return bool(re.match(r"^\d{1,2}\.\s+", _normalize_key(line)))


def _extract_section_lines(text: str | None, start_markers: tuple[str, ...], stop_markers: tuple[str, ...]) -> list[str]:
    if not text:
        return []

    raw_lines = [_clean_topic(line) for line in text.splitlines()]
    lines = [line for line in raw_lines if line]
    start = next(
        (
            index
            for index, line in enumerate(lines)
            if any(marker in _normalize_key(line) for marker in start_markers)
        ),
        None,
    )
    if start is None:
        return []

    section: list[str] = []
    for line in lines[start + 1 :]:
        key = _normalize_key(line)
        if any(marker in key for marker in stop_markers):
            break
        if _is_section_heading(line) and section:
            break
        if "FORMULARIO DE INFORME" in key or "JOTFORM" in key or re.fullmatch(r"\d+\s+DE\s+\d+", key):
            continue
        section.append(line)
    return section


def _is_product_status(line: str) -> bool:
    key = _normalize_key(line)
    return key in {"PUBLICADO", "ENVIADO A REVISION", "ACEPTADO", "EN REVISION", "PUBLICADA", "ENVIADO"}


def _is_product_impact(line: str) -> bool:
    key = _normalize_key(line)
    return "IMPACTO MUNDIAL" in key or "IMPACTO REGIONAL" in key


def _is_url_line(line: str) -> bool:
    key = line.strip().lower()
    return key.startswith(("http", "www.")) or "doi.org" in key or ".com" in key or ".org" in key


def _product_type_from_title(title: str) -> str:
    return "UNCLASSIFIED"


def _product_from_block(block: list[str]) -> dict[str, object] | None:
    useful = [
        line
        for line in block
        if _normalize_key(line)
        not in {"TITULO", "AUTOR 1", "AUTOR 2", "AUTOR 3", "AUTOR 4", "AUTOR 5", "ESTADO", "IMPACTO", "LINK"}
    ]
    if not useful:
        return None

    author_indexes = [index for index, line in enumerate(useful) if _looks_like_person_name(line)]
    if not author_indexes:
        return None

    first_author = author_indexes[0]
    title = _clean_topic(" ".join(_collapse_repeated_words(useful[:first_author])))
    if not title:
        return None

    status = next((line for line in useful if _is_product_status(line)), None)
    impact = next((line for line in useful if _is_product_impact(line)), None)
    link = next((line for line in useful if _is_url_line(line)), None)
    authors = [_title_case_name(useful[index]) for index in author_indexes[:5]]

    return {
        "type": _product_type_from_title(title),
        "title": title,
        "authors": authors,
        "status": status,
        "impact": impact,
        "link": link,
        "source_section": "21",
    }


def _extract_scientific_products(text: str | None) -> list[dict[str, object]]:
    section = _extract_section_lines(
        text,
        ("PRODUCCION CIENTIFICA IMPACTO",),
        (
            "INTERCAMBIO DE CONOCIMIENTO",
            "EVIDENCIAS DE LO REPORTADO",
            "PONENCIAS",
            "CAPITULOS",
            "LIBROS",
        ),
    )
    products: list[dict[str, object]] = []
    block: list[str] = []
    for line in section:
        key = _normalize_key(line)
        if key in {"TITULO", "AUTOR 1", "AUTOR 2", "AUTOR 3", "AUTOR 4", "AUTOR 5", "ESTADO", "IMPACTO", "LINK"}:
            continue
        block.append(line)
        if _is_url_line(line) or (_is_product_status(line) and any(_is_product_impact(item) for item in block)):
            product = _product_from_block(block)
            if product:
                products.append(product)
            block = []

    product = _product_from_block(block)
    if product:
        products.append(product)

    unique: dict[str, dict[str, object]] = {}
    for product in products:
        unique[_normalize_key(product.get("title"))] = product
    return list(unique.values())


def _extract_knowledge_exchange_products(text: str | None) -> list[dict[str, object]]:
    section = _extract_section_lines(
        text,
        ("INTERCAMBIO DE CONOCIMIENTO",),
        ("EVIDENCIAS DE LO REPORTADO", "ANEXOS", "OBSERVACIONES"),
    )
    if not section:
        return []

    ignored = {"FACULTAD INTERVINIENTE", "TIPO", "TITULO"}
    useful = [line for line in section if _normalize_key(line) not in ignored]
    products: list[dict[str, object]] = []
    index = 0
    while index + 2 < len(useful):
        faculty = useful[index]
        kind = useful[index + 1]
        title = useful[index + 2]
        if _normalize_key(kind) in {"CONFERENCIA", "PONENCIA", "CHARLA", "TALLER"}:
            products.append(
                {
                    "type": "PRESENTATION",
                    "title": title,
                    "authors": [],
                    "status": kind,
                    "impact": faculty,
                    "link": None,
                    "source_section": "24",
                }
            )
            index += 3
            continue
        index += 1
    return products


def _extract_all_scientific_products(text: str | None) -> list[dict[str, object]]:
    parsed = parse_progress_report(text)
    if parsed.produccion_cientifica or parsed.intercambios:
        products = list(parsed.produccion_cientifica)
        products.extend(
            {
                "type": "PRESENTATION",
                "title": item.get("title"),
                "authors": [],
                "status": item.get("type"),
                "impact": item.get("faculty"),
                "link": None,
                "source_section": "intercambios",
            }
            for item in parsed.intercambios
        )
        unique: dict[str, dict[str, object]] = {}
        for product in products:
            unique[f"{product.get('type')}|{_normalize_key(product.get('title'))}"] = product
        return list(unique.values())

    products = _extract_scientific_products(text) + _extract_knowledge_exchange_products(text)
    unique: dict[str, dict[str, object]] = {}
    for product in products:
        unique[f"{product.get('type')}|{_normalize_key(product.get('title'))}"] = product
    return list(unique.values())


def _production_counts_from_products(products: list[dict[str, object]]) -> dict[str, int]:
    return {
        "articles": sum(1 for product in products if product.get("type") == "ARTICLE"),
        "books": sum(1 for product in products if product.get("type") == "BOOK"),
        "book_chapters": sum(1 for product in products if product.get("type") == "BOOK_CHAPTER"),
        "presentations": sum(1 for product in products if product.get("type") == "PRESENTATION"),
        "unclassified": sum(1 for product in products if product.get("type") == "UNCLASSIFIED"),
    }


def _extract_group_projects(text: str | None) -> list[dict[str, str]]:
    parsed = parse_progress_report(text)
    if parsed.proyectos_fci:
        return [
            {
                "code": str(project.get("code") or ""),
                "name": str(project.get("name") or ""),
                "director": str(project.get("director") or "Director no detectado"),
                "progress": str(project.get("progress") or ""),
                "status": str(project.get("status") or ""),
                "source_section": str(project.get("source_section") or "proyectos_fci"),
            }
            for project in parsed.proyectos_fci
        ]
    if not text or "CODIGO FCI" not in _normalize_key(text):
        return []

    raw_lines = [_clean_topic(line) for line in text.splitlines()]
    lines = [line for line in raw_lines if line]
    start = next((index for index, line in enumerate(lines) if _normalize_key(line) == "CODIGO FCI"), None)
    if start is None:
        return []

    stop_markers = (
        "FONDOS EXTERNOS",
        "IMPACTOS GENERADOS",
        "PRODUCCION CIENTIFICA",
        "EVIDENCIAS DE LO REPORTADO",
    )
    segment: list[str] = []
    for line in lines[start + 1 :]:
        key = _normalize_key(line)
        if any(marker in key for marker in stop_markers):
            break
        if key in {"PROYECTO", "DIRECTOR", "ESTADO", "% DE AVANCE", "DEL PROYECTO"}:
            continue
        if "FORMULARIO DE INFORME" in key or "JOTFORM" in key or re.fullmatch(r"\d+\s+DE\s+\d+", key):
            continue
        segment.append(line)

    projects: list[dict[str, str]] = []
    code_indexes = [index for index, line in enumerate(segment) if re.fullmatch(r"FCI\d{3,}-\d{4}", _normalize_key(line))]
    for position, index in enumerate(code_indexes):
        next_index = code_indexes[position + 1] if position + 1 < len(code_indexes) else len(segment)
        block = segment[index + 1 : next_index]
        progress_index = next((i for i, line in enumerate(block) if re.fullmatch(r"\d+%", line.strip())), None)
        if progress_index is None:
            continue

        before_progress = block[:progress_index]
        progress = block[progress_index]
        status = block[progress_index + 1] if progress_index + 1 < len(block) else ""
        director = ""
        director_start = len(before_progress)
        for size in range(min(3, len(before_progress)), 0, -1):
            candidate = " ".join(before_progress[-size:])
            if _looks_like_person_name(candidate):
                director = _title_case_name(candidate)
                director_start = len(before_progress) - size
                break

        title = _clean_topic(" ".join(_collapse_repeated_words(before_progress[:director_start])))
        if not title:
            continue

        projects.append(
            {
                "code": segment[index],
                "name": title,
                "director": director or "Director no detectado",
                "progress": progress,
                "status": status,
                "source_section": "proyectos_fci_fallback",
            }
        )
    return projects


def _extract_group_members(text: str | None) -> list[dict[str, str | None]]:
    parsed = parse_progress_report(text)
    if parsed.integrantes_internos:
        return [
            {
                "name": item.get("name"),
                "faculty": item.get("faculty"),
                "career": item.get("career"),
                "source_section": item.get("source_section") or "integrantes_internos",
            }
            for item in parsed.integrantes_internos
        ]
    if not text:
        return []

    raw_lines = [_clean_topic(line) for line in text.splitlines()]
    lines = [line for line in raw_lines if line]
    start = next(
        (
            index
            for index, line in enumerate(lines)
            if "INVESTIGADORES QUE PARTICIPAN EN EL GRUPO" in _normalize_key(line)
            or "INVESTIGADORES QUE PARTICIPAN EN EL PROYECTO" in _normalize_key(line)
        ),
        None,
    )
    if start is None:
        return []

    stop_markers = (
        "INSTITUCIONES E INVESTIGADORES EXTERNOS",
        "INVESTIGADORES QUE SE DESVINCULAN",
        "INVESTIGADORES EXTERNOS QUE PARTICIPAN",
        "ESTUDIANTES QUE PARTICIPAN",
        "RESUMEN DEL PROYECTO",
        "PROYECTOS FCI",
        "CODIGO FCI",
        "ACTIVIDADES",
        "PRODUCCION CIENTIFICA",
    )
    segment: list[str] = []
    for line in lines[start + 1 :]:
        key = _normalize_key(line)
        if any(marker in key for marker in stop_markers):
            break
        if key in {"NOMBRE COMPLETO", "FACULTAD", "CARRERA", "CIENCIAS ADMINISTRATIVAS"}:
            continue
        segment.append(line)

    career_markers = (
        "LICENCIATURA EN",
        "ADMINISTRACION DE EMPRESAS",
        "COMERCIO EXTERIOR",
        "CONTABILIDAD Y AUDITORIA",
        "MERCADOTECNIA",
        "PEDAGOGIA",
    )
    month_markers = (
        "ENERO",
        "FEBRERO",
        "MARZO",
        "ABRIL",
        "MAYO",
        "JUNIO",
        "JULIO",
        "AGOSTO",
        "SEPTIEMBRE",
        "OCTUBRE",
        "NOVIEMBRE",
        "DICIEMBRE",
    )
    members: list[dict[str, str | None]] = []
    index = 0
    while index < len(segment):
        name = segment[index]
        name_cursor = index + 1
        if (
            name_cursor < len(segment)
            and len(_normalize_key(segment[name_cursor]).split()) <= 2
            and _looks_like_person_name(f"{name} {segment[name_cursor]}")
        ):
            name = f"{name} {segment[name_cursor]}"
            name_cursor += 1
        if not _looks_like_person_name(name):
            index += 1
            continue

        faculty: str | None = None
        career_parts: list[str] = []
        cursor = name_cursor
        if cursor < len(segment) and "CIENCIAS ADMINISTRATIVAS" in _normalize_key(segment[cursor]):
            faculty = "CIENCIAS ADMINISTRATIVAS"
            cursor += 1
        elif (
            cursor + 1 < len(segment)
            and _normalize_key(segment[cursor]) == "CIENCIAS"
            and _normalize_key(segment[cursor + 1]) == "ADMINISTRATIVAS"
        ):
            faculty = "CIENCIAS ADMINISTRATIVAS"
            cursor += 2
        elif cursor + 1 < len(segment) and "FILOSOFIA Y LETRAS" in _normalize_key(segment[cursor]):
            faculty = f"{segment[cursor]} {segment[cursor + 1]}"
            cursor += 2

        while cursor < len(segment):
            key = _normalize_key(segment[cursor])
            if _looks_like_person_name(segment[cursor]) and career_parts:
                break
            if any(marker in key for marker in month_markers):
                break
            if any(marker in key for marker in career_markers) or career_parts:
                career_parts.append(segment[cursor])
                cursor += 1
                continue
            break

        clean_name = _clean_person_name(name)
        career_text = " ".join(career_parts)
        clean_career = _clean_career_name(career_text) or _clean_topic(career_text)
        if clean_name and clean_career:
            members.append(
                {
                    "name": clean_name,
                    "faculty": faculty,
                    "career": clean_career,
                    "source_section": "integrantes_internos_fallback",
                }
            )
        index = max(cursor, index + 1)

    unique: dict[str, dict[str, str | None]] = {}
    for member in members:
        unique[_person_key(member["name"])] = member
    return list(unique.values())


def _extract_external_researchers(text: str | None) -> list[dict[str, str | None]]:
    parsed = parse_progress_report(text)
    if parsed.integrantes_externos:
        return [
            {
                "name": item.get("name"),
                "institution": item.get("institution"),
                "is_external": "true",
                "source_section": item.get("source_section") or "integrantes_externos",
            }
            for item in parsed.integrantes_externos
        ]
    if not text:
        return []

    raw_lines = [_clean_topic(line) for line in text.splitlines()]
    lines = [line for line in raw_lines if line]
    start = next(
        (
            index
            for index, line in enumerate(lines)
            if "INSTITUCIONES E INVESTIGADORES EXTERNOS" in _normalize_key(line)
            or "INVESTIGADORES EXTERNOS QUE PARTICIPAN" in _normalize_key(line)
        ),
        None,
    )
    if start is None:
        return []

    segment: list[str] = []
    for line in lines[start + 1 :]:
        key = _normalize_key(line)
        if (
            key.startswith("14")
            or key.startswith("17")
            or "RESUMEN DEL GRUPO" in key
            or "RESUMEN DEL PROYECTO" in key
            or "ESTUDIANTES QUE PARTICIPAN" in key
            or "SECCION 2" in key
        ):
            break
        if key in {"INSTITUCION EXTERNA", "INSTITUCIÓN EXTERNA", "INVESTIGADOR", "NOMBRE COMPLETO"}:
            continue
        if _is_pdf_noise_line(line):
            continue
        segment.append(line)

    researchers: list[dict[str, str | None]] = []
    index = 0
    while index < len(segment):
        if _looks_like_person_name(segment[index]):
            name = segment[index]
            institution_parts: list[str] = []
            cursor = index + 1
            while cursor < len(segment) and not _looks_like_person_name(segment[cursor]):
                candidate_key = _normalize_key(segment[cursor])
                if candidate_key.startswith(("17", "18")) or "ESTUDIANTES QUE PARTICIPAN" in candidate_key:
                    break
                if not _is_pdf_noise_line(segment[cursor]):
                    institution_parts.append(segment[cursor])
                cursor += 1
            clean_name = _clean_person_name(name)
            clean_institution = _clean_external_institution(" ".join(institution_parts))
            if clean_name and clean_institution:
                researchers.append(
                    {
                        "name": clean_name,
                        "institution": clean_institution,
                        "is_external": "true",
                        "source_section": "integrantes_externos_fallback",
                    }
                )
            index = max(cursor, index + 1)
            continue

        if index + 1 < len(segment) and _looks_like_person_name(segment[index + 1]):
            institution = segment[index]
            name = segment[index + 1]
            clean_name = _clean_person_name(name)
            clean_institution = _clean_external_institution(institution)
            if clean_name and clean_institution:
                researchers.append(
                    {
                        "name": clean_name,
                        "institution": clean_institution,
                        "is_external": "true",
                        "source_section": "integrantes_externos_fallback",
                    }
                )
            index += 2
            continue

        index += 1

    unique: dict[str, dict[str, str | None]] = {}
    for researcher in researchers:
        key = _person_key(researcher["name"])
        existing = unique.get(key)
        if not existing:
            unique[key] = researcher
            continue
        current_institution = existing.get("institution")
        new_institution = researcher.get("institution")
        if current_institution and new_institution and _normalize_key(current_institution) != _normalize_key(new_institution):
            existing["institution"] = f"{current_institution} / {new_institution}"
    return list(unique.values())


class ImportService:
    def __init__(self, db: Session):
        self.db = db

    async def import_progress_pdf(
        self,
        file: UploadFile,
        imported_by: str | None,
        source_path: str | None = None,
    ) -> ImportResult:
        content = await file.read()
        return await self._import_progress_pdf_content(
            content=content,
            imported_by=imported_by,
            source_path=source_path,
            filename=file.filename,
        )

    async def import_progress_pdf_batch(
        self,
        payload: DropboxProgressPdfBatchRequest,
        imported_by: str | None,
    ) -> list[ImportResult]:
        if not payload.files:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail="No se recibieron archivos de Dropbox para importar.",
            )

        results: list[ImportResult] = []
        for item in payload.files:
            filename = item.name or PurePosixPath(item.path_lower).name
            source_path = item.path_display or item.path_lower
            dropbox_metadata = self._dropbox_metadata(item, filename, source_path)
            if self._dropbox_revision_exists(dropbox_metadata):
                results.append(
                    ImportResult(
                        job_id=0,
                        source_type="PROGRESS_PDF",
                        filename=filename,
                        status="SKIPPED",
                        imported_rows=0,
                        summary="Esta revision de Dropbox ya fue procesada; se omite para evitar duplicados.",
                    )
                )
                continue
            try:
                content = await self._download_dropbox_file(item)
                result = await self._import_progress_pdf_content(
                    content=content,
                    imported_by=imported_by,
                    source_path=source_path,
                    filename=filename,
                    dropbox_metadata=dropbox_metadata,
                )
                results.append(result)
            except HTTPException as exc:
                self.db.rollback()
                results.append(
                    ImportResult(
                        job_id=0,
                        source_type="PROGRESS_PDF",
                        filename=filename,
                        status="ERROR",
                        imported_rows=0,
                        summary=f"No se pudo procesar el archivo: {exc.detail}",
                    )
                )
            except Exception as exc:
                self.db.rollback()
                results.append(
                    ImportResult(
                        job_id=0,
                        source_type="PROGRESS_PDF",
                        filename=filename,
                        status="ERROR",
                        imported_rows=0,
                        summary=f"Error inesperado al procesar el archivo: {exc}",
                    )
                )

        return results

    async def _import_progress_pdf_content(
        self,
        content: bytes,
        imported_by: str | None,
        source_path: str | None = None,
        filename: str | None = None,
        dropbox_metadata: dict | None = None,
        existing_job: ImportJob | None = None,
    ) -> ImportResult:
        if not self._is_pdf_content(content):
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail="Solo se permiten archivos PDF.",
            )

        resolved_filename = self._resolve_pdf_filename(filename, source_path)
        job = existing_job or self._create_job("PROGRESS_PDF", resolved_filename, imported_by)
        if dropbox_metadata:
            job.source_identifier = (
                dropbox_metadata.get("dropbox_id")
                or dropbox_metadata.get("path_lower")
                or source_path
            )
            job.source_rev = dropbox_metadata.get("rev")
            job.source_fingerprint = dropbox_fingerprint(dropbox_metadata) or None
            job.document_key = dropbox_document_key(dropbox_metadata) or None
            job.is_current = False
        job.status = "PROCESSING"
        job.current_step = "validating_pdf"
        self.db.flush()
        processing_started = perf_counter()
        extraction_started = perf_counter()
        job.current_step = "extracting_text"
        self.db.flush()
        extraction = self._extract_text_from_pdf(content)
        job.text_extraction_ms = int((perf_counter() - extraction_started) * 1000)
        job.ocr_ms = extraction.ocr_ms
        job.page_count = extraction.page_count
        job.used_ocr = extraction.used_ocr
        digital_pages = sum(1 for item in extraction.page_logs if item.method == "PyMuPDF" and item.chars > 0)
        ocr_pages = sum(1 for item in extraction.page_logs if item.method == "Tesseract")
        job.extraction_method = "mixed" if digital_pages and ocr_pages else ("ocr" if ocr_pages else "digital_text")
        extracted_text = extraction.text
        extraction_provider = extraction.provider
        extraction_issue = extraction.issue
        extracted_text = _sanitize_db_text(extracted_text) or ""
        extraction_issue = _sanitize_db_text(extraction_issue)
        job.current_step = "classifying_document"
        self.db.flush()
        classification = classify_document(extracted_text, resolved_filename)
        parser_started = perf_counter()
        job.current_step = "parsing_payload"
        self.db.flush()
        parsed_payload = self._build_pdf_parse_payload(
            extracted_text=extracted_text,
            filename=resolved_filename,
            source_path=source_path,
            extraction=extraction,
            dropbox_metadata=dropbox_metadata,
            classification=classification,
        )
        job.parser_ms = int((perf_counter() - parser_started) * 1000)
        persistence_started = perf_counter()
        job.current_step = "persisting_trace"
        self.db.flush()

        if not extracted_text.strip():
            review_notes = (
                extraction_issue
                or "El PDF fue recibido correctamente, pero no contiene texto extraible en esta primera pasada automatica."
            )
            trace = ImportedOcrTrace(
                import_job_id=job.id,
                progress_report_id=None,
                source_filename=_normalize_text(resolved_filename),
                source_path=_normalize_text(source_path) or None,
                ocr_provider=extraction_provider,
                extracted_text=None,
                parsed_payload=parsed_payload,
                confidence_score=None,
                review_status="OMITIDO_SIN_TEXTO",
                review_notes=review_notes,
            )
            self.db.add(trace)
            job.persistence_ms = int((perf_counter() - persistence_started) * 1000)
            summary = "Se recibio el PDF, pero no fue posible extraer texto automaticamente. El archivo se omite sin detener el lote."
            status_value = "SKIPPED" if classification.status == "ignored" else "REQUIRES_REVIEW"
            return self._finalize_job(job, 0, summary, status_value=status_value)

        if classification.status != "accepted":
            trace = ImportedOcrTrace(
                import_job_id=job.id,
                progress_report_id=None,
                source_filename=_normalize_text(resolved_filename),
                source_path=_normalize_text(source_path) or None,
                ocr_provider=extraction_provider,
                extracted_text=extracted_text,
                parsed_payload=parsed_payload,
                confidence_score=classification.score / 100,
                review_status=(
                    "OMITIDO_NO_INFORME"
                    if classification.status == "ignored"
                    else "REQUIERE_REVISION_TIPO_DOCUMENTO"
                ),
                review_notes="; ".join(classification.warnings) or None,
            )
            self.db.add(trace)
            job.persistence_ms = int((perf_counter() - persistence_started) * 1000)
            summary = (
                "El PDF fue recibido, pero no parece ser un informe parcial de GI; se omitio sin contaminar los KPIs."
                if classification.status == "ignored"
                else "El PDF fue recibido, pero necesita revision porque no hay senales suficientes de informe parcial."
            )
            status_value = "SKIPPED" if classification.status == "ignored" else "REQUIRES_REVIEW"
            return self._finalize_job(job, 0, summary, status_value=status_value)

        trace = ImportedOcrTrace(
            import_job_id=job.id,
            progress_report_id=None,
            source_filename=_normalize_text(resolved_filename),
            source_path=_normalize_text(source_path) or None,
            ocr_provider=extraction_provider,
            extracted_text=extracted_text,
            parsed_payload=parsed_payload,
            confidence_score=classification.score / 100,
            review_status="PENDIENTE_REVISION",
        )
        self.db.add(trace)
        self.db.flush()

        progress = self._build_progress_record_from_pdf_text(
            job_id=job.id,
            filename=resolved_filename,
            source_path=source_path,
            extracted_text=extracted_text,
        ) if not self._progress_exists_for_source(resolved_filename, source_path, dropbox_metadata) else None
        if progress:
            self.db.add(progress)
            self.db.flush()
            trace.progress_report_id = progress.id
            trace.review_status, trace.review_notes = self._accepted_pdf_review_state(
                classification_score=classification.score,
                parsed_payload=parsed_payload,
                progress=progress,
            )

        persisted_counts = self._persist_normalized_from_payload(parsed_payload, progress)
        parsed_payload["persisted_counts"] = persisted_counts
        trace.parsed_payload = parsed_payload
        flag_modified(trace, "parsed_payload")
        imported_rows = 1 if progress else 0
        job.persistence_ms = int((perf_counter() - persistence_started) * 1000)
        job.duration_ms = int((perf_counter() - processing_started) * 1000)
        summary = (
            "Se recibio el PDF de avance, se extrajo su texto y se registro una investigacion preliminar para analisis."
            if progress
            else "Se recibio el PDF de avance y se extrajo su texto para revision y posterior parseo."
        )
        return self._finalize_job(job, imported_rows, summary)

    async def _download_dropbox_file(self, item: DropboxProgressPdfItem) -> bytes:
        access_token = await self._get_dropbox_access_token()
        headers = {
            "Authorization": f"Bearer {access_token}",
            "Dropbox-API-Arg": json.dumps({"path": item.path_lower}),
        }

        async with httpx.AsyncClient(timeout=120.0) as client:
            response = await client.post(
                "https://content.dropboxapi.com/2/files/download",
                headers=headers,
            )

        if response.status_code >= 400:
            raise HTTPException(
                status_code=status.HTTP_502_BAD_GATEWAY,
                detail=(
                    f"Dropbox devolvio un error al descargar '{item.name or item.path_lower}': "
                    f"{response.status_code}"
                ),
            )

        return response.content

    async def _get_dropbox_access_token(self) -> str:
        if not settings.dropbox_client_id or not settings.dropbox_client_secret or not settings.dropbox_refresh_token:
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail="Faltan las credenciales de Dropbox en el backend.",
            )

        async with httpx.AsyncClient(timeout=30.0) as client:
            response = await client.post(
                "https://api.dropboxapi.com/oauth2/token",
                data={
                    "grant_type": "refresh_token",
                    "refresh_token": settings.dropbox_refresh_token,
                    "client_id": settings.dropbox_client_id,
                    "client_secret": settings.dropbox_client_secret,
                },
                headers={"Content-Type": "application/x-www-form-urlencoded"},
            )

        if response.status_code >= 400:
            raise HTTPException(
                status_code=status.HTTP_502_BAD_GATEWAY,
                detail="No se pudo renovar el token de Dropbox.",
            )

        payload = response.json()
        access_token = payload.get("access_token")
        if not access_token:
            raise HTTPException(
                status_code=status.HTTP_502_BAD_GATEWAY,
                detail="Dropbox no devolvio un access token valido.",
            )
        return access_token

    async def import_research_base(self, file: UploadFile, imported_by: str | None) -> ImportResult:
        rows = self._load_rows(file)
        required = [
            "FACULTAD",
            "CEDULA",
            "APELLIDOS Y NOMBRES",
            "GENERO",
            "NOMBRE DEL PROYECTO",
            "CODIGO",
            "ANO",
            "ROL",
            "DEDICACION",
        ]
        self._validate_columns(rows["headers"], required)

        job = self._create_job("RESEARCH_BASE", file.filename or "research_base.xlsx", imported_by)
        imported = 0
        for item in rows["records"]:
            record = ImportedResearchRecord(
                import_job_id=job.id,
                faculty_name=_normalize_text(item.get("FACULTAD")) or None,
                national_id=_normalize_text(item.get("CEDULA")),
                full_name=_normalize_text(item.get("APELLIDOS Y NOMBRES")),
                gender=_normalize_text(item.get("GENERO")) or None,
                project_name=_normalize_text(item.get("NOMBRE DEL PROYECTO")) or None,
                project_code=_normalize_text(item.get("CODIGO")) or None,
                year_label=_normalize_text(item.get("ANO")) or None,
                role_name=_normalize_text(item.get("ROL")) or None,
                dedication=_normalize_text(item.get("DEDICACION")) or None,
                notes=_normalize_text(item.get("OBSERVACION")) or None,
            )
            self.db.add(record)
            imported += 1

        summary = (
            f"Se importaron {imported} registros de la base de investigadores para la Facultad de Ciencias Administrativas."
        )
        return self._finalize_job(job, imported, summary)

    async def import_project_participants(
        self,
        file: UploadFile,
        imported_by: str | None,
    ) -> ImportResult:
        rows = self._load_rows(file)
        required = [
            "CODIGO_IES",
            "CODIGO",
            "NOMBRE",
            "TIPO_PROYECTO",
            "TIPO_PARTICIPANTE",
            "IDENTIFICACION_CODIGO",
            "HORAS",
            "GRUPO_INVESTIGACION",
        ]
        self._validate_columns(rows["headers"], required)

        job = self._create_job(
            "PROJECT_PARTICIPANTS",
            file.filename or "project_participants.xlsx",
            imported_by,
        )
        imported = 0
        for item in rows["records"]:
            record = ImportedProjectParticipant(
                import_job_id=job.id,
                ies_code=_normalize_text(item.get("CODIGO_IES")) or None,
                project_code=_normalize_text(item.get("CODIGO")),
                project_name=_normalize_text(item.get("NOMBRE")),
                project_type=_normalize_text(item.get("TIPO_PROYECTO")) or None,
                participant_type=_normalize_text(item.get("TIPO_PARTICIPANTE")) or None,
                participant_identifier=_normalize_text(item.get("IDENTIFICACION_CODIGO")),
                assigned_hours=_normalize_int(item.get("HORAS"), default=0),
                research_group=_normalize_text(item.get("GRUPO_INVESTIGACION")) or None,
            )
            self.db.add(record)
            imported += 1

        summary = f"Se importaron {imported} filas de proyectos y participantes."
        return self._finalize_job(job, imported, summary)

    async def import_progress_report(self, file: UploadFile, imported_by: str | None) -> ImportResult:
        rows = self._load_rows(file)
        required = [
            "CARRERA",
            "ANO",
            "CICLO",
            "DOCENTE",
            "ARTICULOS",
            "LIBROS",
            "CAPITULOS",
            "PONENCIAS",
            "PROYECTOS",
        ]
        self._validate_columns(rows["headers"], required)

        job = self._create_job("PROGRESS_REPORT", file.filename or "progress_report.xlsx", imported_by)
        imported = 0
        for item in rows["records"]:
            record = ImportedProgressReport(
                import_job_id=job.id,
                career_name=_normalize_text(item.get("CARRERA")) or None,
                year_label=_normalize_text(item.get("ANO")),
                cycle=_normalize_int(item.get("CICLO"), default=1),
                teacher_identifier=_normalize_text(item.get("CEDULA")) or None,
                teacher_name=_normalize_text(item.get("DOCENTE")),
                articles=_normalize_int(item.get("ARTICULOS")),
                books=_normalize_int(item.get("LIBROS")),
                book_chapters=_normalize_int(item.get("CAPITULOS")),
                presentations=_normalize_int(item.get("PONENCIAS")),
                projects=_normalize_int(item.get("PROYECTOS")),
                notes=_normalize_text(item.get("OBSERVACIONES")) or None,
            )
            self.db.add(record)
            imported += 1

        summary = f"Se importaron {imported} filas de avance para comparacion con POA."
        return self._finalize_job(job, imported, summary)

    def import_progress_json(
        self,
        payload: ProgressJsonImportRequest,
        imported_by: str | None,
    ) -> ImportResult:
        if not payload.records:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail="No se recibieron registros de avance para importar.",
            )

        job = self._create_job("PROGRESS_JSON", payload.source_filename, imported_by)
        imported = 0

        for item in payload.records:
            record = ImportedProgressReport(
                import_job_id=job.id,
                career_name=_normalize_text(item.career_name) or None,
                year_label=_normalize_text(item.year_label),
                cycle=item.cycle,
                teacher_identifier=_normalize_text(item.teacher_identifier) or None,
                teacher_name=_normalize_text(item.teacher_name),
                articles=item.articles,
                books=item.books,
                book_chapters=item.book_chapters,
                presentations=item.presentations,
                projects=item.projects,
                notes=_normalize_text(item.notes) or None,
            )
            self.db.add(record)
            self.db.flush()

            trace = ImportedOcrTrace(
                import_job_id=job.id,
                progress_report_id=record.id,
                source_filename=_normalize_text(payload.source_filename),
                source_path=_normalize_text(payload.source_path) or None,
                ocr_provider=_normalize_text(payload.ocr_provider) or None,
                extracted_text=_sanitize_db_text(_normalize_text(payload.extracted_text) or None),
                parsed_payload=item.model_dump(),
                confidence_score=payload.confidence_score,
                review_status="PENDIENTE_REVISION",
            )
            self.db.add(trace)
            imported += 1

        summary = f"Se importaron {imported} registros de avance desde OCR para revision."
        return self._finalize_job(job, imported, summary)

    async def import_poa(self, file: UploadFile, imported_by: str | None) -> ImportResult:
        rows = self._load_rows(file)
        required = ["CARRERA", "ANO", "METRICA", "VALOR_PLANIFICADO"]
        self._validate_columns(rows["headers"], required)

        job = self._create_job("POA", file.filename or "poa.xlsx", imported_by)
        imported = 0
        career_lookup = {_normalize_header(career.name): career for career in self.db.query(Career).all()}
        metric_lookup = {_normalize_header(metric.value): metric for metric in GoalMetric}

        for item in rows["records"]:
            career_name = _normalize_header(item.get("CARRERA"))
            career = career_lookup.get(career_name)
            if not career:
                raise HTTPException(
                    status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                    detail=f"La carrera '{career_name}' no existe en el sistema.",
                )

            metric_name = _normalize_header(item.get("METRICA"))
            metric = metric_lookup.get(metric_name)
            if not metric:
                raise HTTPException(
                    status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                    detail=f"La metrica '{metric_name}' no esta soportada.",
                )

            year_label = _normalize_text(item.get("ANO"))
            planned_value = _normalize_int(item.get("VALOR_PLANIFICADO"))

            goal = (
                self.db.query(AnnualGoal)
                .filter(
                    AnnualGoal.career_id == career.id,
                    AnnualGoal.year_label == year_label,
                    AnnualGoal.metric == metric,
                )
                .first()
            )
            if goal:
                goal.planned_value = planned_value
            else:
                self.db.add(
                    AnnualGoal(
                        career_id=career.id,
                        year_label=year_label,
                        metric=metric,
                        planned_value=planned_value,
                    )
                )
            imported += 1

        summary = f"Se importaron {imported} metas POA."
        return self._finalize_job(job, imported, summary)

    def list_ocr_traces(self) -> list[ImportedOcrTrace]:
        return self.db.query(ImportedOcrTrace).order_by(ImportedOcrTrace.created_at.desc()).all()

    def sync_ocr_traces_to_progress(self) -> dict[str, int]:
        self._clear_generated_progress_records()
        traces = (
            self.db.query(ImportedOcrTrace)
            .filter(
                ImportedOcrTrace.progress_report_id.is_(None),
                ImportedOcrTrace.extracted_text.is_not(None),
            )
            .order_by(ImportedOcrTrace.created_at.desc())
            .all()
        )

        created = 0
        skipped = 0
        seen_sources: set[str] = set()
        for trace in traces:
            source_key = self._ocr_source_key(trace.source_filename, trace.source_path)
            if source_key in seen_sources:
                skipped += 1
                continue
            seen_sources.add(source_key)

            progress = self._build_progress_record_from_pdf_text(
                job_id=trace.import_job_id,
                filename=trace.source_filename,
                source_path=trace.source_path,
                extracted_text=trace.extracted_text or "",
            )
            if not progress:
                skipped += 1
                continue

            self.db.add(progress)
            self.db.flush()
            trace.progress_report_id = progress.id
            created += 1

        self.db.commit()
        return {"created": created, "skipped": skipped}

    def backfill_missing_external_researchers(self) -> dict[str, int]:
        traces = (
            self.db.query(ImportedOcrTrace)
            .filter(ImportedOcrTrace.extracted_text.is_not(None))
            .order_by(ImportedOcrTrace.id.desc())
            .all()
        )
        scanned = 0
        updated = 0
        still_empty = 0
        skipped = 0

        for trace in traces:
            scanned += 1
            payload = dict(trace.parsed_payload or {})
            current = payload.get("integrantes_externos")
            if isinstance(current, list) and current:
                skipped += 1
                continue

            parsed = parse_progress_report(trace.extracted_text or "").to_public_dict()
            external_researchers = parsed.get("integrantes_externos") or []
            if not external_researchers:
                still_empty += 1
                continue

            payload["integrantes_externos"] = external_researchers
            payload["logs"] = parsed.get("logs", payload.get("logs", []))
            payload["extraction_log"] = parsed.get("extraction_log", payload.get("extraction_log", []))
            payload["estado"] = parsed.get("estado", payload.get("estado", "accepted"))
            trace.parsed_payload = payload
            updated += 1

        self.db.commit()
        return {
            "scanned": scanned,
            "updated": updated,
            "skipped": skipped,
            "still_empty": still_empty,
        }

    def _rebuild_trace_payload_from_text(self, trace: ImportedOcrTrace) -> dict:
        current_payload = dict(trace.parsed_payload or {})
        parsed = parse_progress_report(trace.extracted_text or "").to_public_dict()
        document = current_payload.get("document") or {
            "filename": trace.source_filename,
            "source_path": trace.source_path,
            "source_key": self._ocr_source_key(trace.source_filename, trace.source_path),
            "type": "progress_report_pdf",
            "status": "accepted",
            "score": int((trace.confidence_score or 1) * 100),
        }
        page_extraction_log = current_payload.get("page_extraction_log") or []
        parser_logs = parsed.get("logs") or []
        parsed["document"] = document
        parsed["warnings"] = current_payload.get("warnings") or []
        parsed["page_extraction_log"] = page_extraction_log
        parsed["extraction_log"] = [
            *[
                {
                    "field": "pdf.page",
                    "value": item.get("chars"),
                    "page": item.get("page"),
                    "section": None,
                    "method": item.get("method"),
                    "score": 100 if item.get("chars") else 0,
                    "requires_review": item.get("requires_review", False),
                    "warning": item.get("warning"),
                }
                for item in page_extraction_log
                if isinstance(item, dict)
            ],
            *[{**log, "section": log.get("table")} for log in parser_logs if isinstance(log, dict)],
        ]
        return parsed

    def _period_ids_for_import_jobs(self, job_ids: list[int]) -> list[int]:
        if not job_ids:
            return []
        progresses = (
            self.db.query(ImportedProgressReport.year_label, ImportedProgressReport.cycle)
            .filter(ImportedProgressReport.import_job_id.in_(job_ids))
            .all()
        )
        period_ids: set[int] = set()
        for year_label, cycle in progresses:
            period = (
                self.db.query(AcademicPeriod.id)
                .filter(AcademicPeriod.year_label == year_label, AcademicPeriod.cycle == cycle)
                .first()
            )
            if period:
                period_ids.add(int(period[0]))
        return sorted(period_ids)

    @staticmethod
    def _ids(query) -> set[int]:
        return {int(row[0]) for row in query.all()}

    def _normalization_scope_ids_for_batch(
        self,
        batch_id: int,
        job_ids: list[int],
        period_ids: list[int],
    ) -> dict[str, set[int]]:
        teacher_ids = self._ids(self.db.query(Teacher.id).filter(Teacher.import_batch_id == batch_id))
        external_ids = self._ids(self.db.query(ExternalResearcher.id).filter(ExternalResearcher.import_batch_id == batch_id))
        production_ids = self._ids(
            self.db.query(ScientificProduction.id).filter(ScientificProduction.import_batch_id == batch_id)
        )
        research_entity_ids = self._ids(
            self.db.query(ResearchEntity.id).filter(ResearchEntity.import_batch_id == batch_id)
        )
        project_ids = self._ids(self.db.query(ResearchProject.id).filter(ResearchProject.import_batch_id == batch_id))

        if job_ids:
            teacher_ids.update(self._ids(self.db.query(Teacher.id).filter(Teacher.import_job_id.in_(job_ids))))
            external_ids.update(self._ids(self.db.query(ExternalResearcher.id).filter(ExternalResearcher.import_job_id.in_(job_ids))))
            production_ids.update(
                self._ids(self.db.query(ScientificProduction.id).filter(ScientificProduction.import_job_id.in_(job_ids)))
            )
            research_entity_ids.update(
                self._ids(self.db.query(ResearchEntity.id).filter(ResearchEntity.import_job_id.in_(job_ids)))
            )
            project_ids.update(self._ids(self.db.query(ResearchProject.id).filter(ResearchProject.import_job_id.in_(job_ids))))

        # Legacy compatibility: records created before traceability had no batch/job id.
        teacher_ids.update(
            self._ids(self.db.query(Teacher.id).filter(Teacher.institutional_email.like("imported-%@local.import")))
        )
        if period_ids:
            external_ids.update(
                self._ids(
                    self.db.query(ExternalResearcher.id).filter(
                        ExternalResearcher.import_batch_id.is_(None),
                        ExternalResearcher.period_id.in_(period_ids),
                    )
                )
            )
            project_ids.update(
                self._ids(
                    self.db.query(ResearchProject.id).filter(
                        ResearchProject.import_batch_id.is_(None),
                        ResearchProject.period_id.in_(period_ids),
                    )
                )
            )
            production_ids.update(
                self._ids(
                    self.db.query(ScientificProduction.id).filter(
                        ScientificProduction.import_batch_id.is_(None),
                        ScientificProduction.period_id.in_(period_ids),
                    )
                )
            )
            research_entity_ids.update(
                self._ids(
                    self.db.query(ResearchEntity.id).filter(
                        ResearchEntity.import_batch_id.is_(None),
                        ResearchEntity.period_id.in_(period_ids),
                    )
                )
            )

        if teacher_ids:
            production_ids.update(
                self._ids(self.db.query(ScientificProduction.id).filter(ScientificProduction.teacher_id.in_(teacher_ids)))
            )
        if research_entity_ids:
            production_ids.update(
                self._ids(
                    self.db.query(ScientificProduction.id).filter(
                        ScientificProduction.research_entity_id.in_(research_entity_ids)
                    )
                )
            )

        return {
            "teachers": teacher_ids,
            "external_researchers": external_ids,
            "scientific_productions": production_ids,
            "research_entities": research_entity_ids,
            "research_projects": project_ids,
        }

    @staticmethod
    def _normalization_scope_counts(scope: dict[str, set[int]]) -> dict[str, int]:
        return {name: len(ids) for name, ids in scope.items()}

    def _clear_normalized_for_batch(self, batch_id: int, job_ids: list[int], period_ids: list[int]) -> dict[str, int]:
        scope = self._normalization_scope_ids_for_batch(batch_id, job_ids, period_ids)
        deleted = {name: 0 for name in scope}
        deleted["project_teachers"] = 0
        deleted["person_roles"] = 0
        deleted["scientific_production_authors"] = 0
        deleted["review_items"] = 0
        deleted["normalization_audits"] = 0

        if job_ids:
            deleted["person_roles"] += self.db.query(PersonRole).filter(
                PersonRole.import_job_id.in_(job_ids)
            ).delete(synchronize_session=False)
            deleted["review_items"] = self.db.query(ImportReviewItem).filter(
                ImportReviewItem.import_job_id.in_(job_ids)
            ).delete(synchronize_session=False)
            deleted["normalization_audits"] += self.db.query(ImportNormalizationAudit).filter(
                ImportNormalizationAudit.import_job_id.in_(job_ids)
            ).delete(synchronize_session=False)
        deleted["normalization_audits"] += self.db.query(ImportNormalizationAudit).filter(
            ImportNormalizationAudit.import_batch_id == batch_id
        ).delete(synchronize_session=False)
        deleted["person_roles"] += self.db.query(PersonRole).filter(
            PersonRole.import_batch_id == batch_id
        ).delete(synchronize_session=False)

        project_ids = scope["research_projects"]
        teacher_ids = scope["teachers"]
        external_ids = scope["external_researchers"]
        production_ids = scope["scientific_productions"]
        research_entity_ids = scope["research_entities"]
        if teacher_ids:
            deleted["person_roles"] += self.db.query(PersonRole).filter(
                PersonRole.teacher_id.in_(teacher_ids)
            ).delete(synchronize_session=False)
        if external_ids:
            deleted["person_roles"] += self.db.query(PersonRole).filter(
                PersonRole.external_researcher_id.in_(external_ids)
            ).delete(synchronize_session=False)
        if project_ids:
            deleted["person_roles"] += self.db.query(PersonRole).filter(
                PersonRole.research_project_id.in_(project_ids)
            ).delete(synchronize_session=False)
        if production_ids:
            deleted["person_roles"] += self.db.query(PersonRole).filter(
                PersonRole.scientific_production_id.in_(production_ids)
            ).delete(synchronize_session=False)
        if research_entity_ids:
            deleted["person_roles"] += self.db.query(PersonRole).filter(
                PersonRole.research_entity_id.in_(research_entity_ids)
            ).delete(synchronize_session=False)
        if project_ids:
            deleted["project_teachers"] += self.db.query(ProjectTeacher).filter(
                ProjectTeacher.project_id.in_(project_ids)
            ).delete(synchronize_session=False)
        if teacher_ids:
            deleted["project_teachers"] += self.db.query(ProjectTeacher).filter(
                ProjectTeacher.teacher_id.in_(teacher_ids)
            ).delete(synchronize_session=False)
        if production_ids:
            deleted["scientific_production_authors"] += self.db.query(ScientificProductionAuthor).filter(
                ScientificProductionAuthor.production_id.in_(production_ids)
            ).delete(synchronize_session=False)
            deleted["scientific_production_authors"] += self.db.query(ScientificProductionAuthor).filter(
                ScientificProductionAuthor.import_batch_id == batch_id
            ).delete(synchronize_session=False)
            if job_ids:
                deleted["scientific_production_authors"] += self.db.query(ScientificProductionAuthor).filter(
                    ScientificProductionAuthor.import_job_id.in_(job_ids)
                ).delete(synchronize_session=False)
            if research_entity_ids:
                deleted["scientific_production_authors"] += self.db.query(ScientificProductionAuthor).filter(
                    ScientificProductionAuthor.research_entity_id.in_(research_entity_ids)
                ).delete(synchronize_session=False)
            deleted["scientific_productions"] = self.db.query(ScientificProduction).filter(
                ScientificProduction.id.in_(production_ids)
            ).delete(synchronize_session=False)
        if research_entity_ids:
            deleted["research_entities"] = self.db.query(ResearchEntity).filter(
                ResearchEntity.id.in_(research_entity_ids)
            ).delete(synchronize_session=False)
        if project_ids:
            deleted["research_projects"] = self.db.query(ResearchProject).filter(
                ResearchProject.id.in_(project_ids)
            ).delete(synchronize_session=False)
        if scope["external_researchers"]:
            deleted["external_researchers"] = self.db.query(ExternalResearcher).filter(
                ExternalResearcher.id.in_(scope["external_researchers"])
            ).delete(synchronize_session=False)
        if teacher_ids:
            deleted["teachers"] = self.db.query(Teacher).filter(Teacher.id.in_(teacher_ids)).delete(
                synchronize_session=False
            )
        self.db.flush()
        return deleted

    def _normalization_audit_summary(self, batch_id: int) -> dict[str, dict[str, int]]:
        rows = (
            self.db.query(ImportNormalizationAudit)
            .filter(ImportNormalizationAudit.import_batch_id == batch_id)
            .order_by(ImportNormalizationAudit.entity_type.asc(), ImportNormalizationAudit.action.asc())
            .all()
        )
        summary: dict[str, dict[str, int]] = {}
        for row in rows:
            by_action = summary.setdefault(row.entity_type, {})
            by_action[row.action] = by_action.get(row.action, 0) + 1
        return summary

    def clear_generated_progress_records(self) -> dict[str, int]:
        deleted = self._clear_generated_progress_records()
        self.db.commit()
        return {"deleted": deleted}

    def _clear_generated_progress_records(self) -> int:
        generated_ids = [
            row[0]
            for row in self.db.query(ImportedProgressReport.id)
            .filter(
                ImportedProgressReport.notes.like(f"{GENERATED_PROGRESS_NOTE}%")
                | ImportedProgressReport.notes.like(f"{LEGACY_GENERATED_PROGRESS_NOTE}%")
            )
            .all()
        ]
        if not generated_ids:
            return 0

        self.db.query(ImportedOcrTrace).filter(
            ImportedOcrTrace.progress_report_id.in_(generated_ids)
        ).update({ImportedOcrTrace.progress_report_id: None}, synchronize_session=False)
        self.db.query(ImportedProgressReport).filter(
            ImportedProgressReport.id.in_(generated_ids)
        ).delete(synchronize_session=False)
        self.db.flush()
        return len(generated_ids)

    @staticmethod
    def _dropbox_metadata(
        item: DropboxProgressPdfItem,
        filename: str | None,
        source_path: str | None,
    ) -> dict:
        source_key = ImportService._ocr_source_key(filename, source_path)
        return {
            "dropbox_id": _normalize_text(item.id) or None,
            "rev": _normalize_text(item.rev) or None,
            "name": _normalize_text(item.name or filename) or None,
            "path_lower": _normalize_text(item.path_lower) or None,
            "path_display": _normalize_text(item.path_display or source_path) or None,
            "client_modified": _normalize_text(item.client_modified) or None,
            "server_modified": _normalize_text(item.server_modified) or None,
            "size": item.size,
            "content_hash": _normalize_text(item.content_hash) or None,
            "source_key": source_key,
        }

    def _dropbox_revision_exists(self, metadata: dict | None) -> bool:
        if not metadata:
            return False
        dropbox_id = metadata.get("dropbox_id")
        rev = metadata.get("rev")
        source_key = metadata.get("source_key")
        document_key = dropbox_document_key(metadata)
        if dropbox_id and rev:
            existing_job = (
                self.db.query(ImportJob.id)
                .filter(
                    ImportJob.document_key == document_key,
                    ImportJob.source_rev == rev,
                )
                .first()
            )
            if existing_job:
                return True
        elif document_key:
            existing_job = (
                self.db.query(ImportJob.id)
                .filter(ImportJob.document_key == document_key)
                .first()
            )
            if existing_job:
                return True
        traces = self.db.query(ImportedOcrTrace).all()
        for trace in traces:
            payload = trace.parsed_payload or {}
            document = payload.get("document") or {}
            if dropbox_id and rev:
                if (
                    trace.progress_report_id
                    and document.get("dropbox_id") == dropbox_id
                    and document.get("rev") == rev
                ):
                    return True
                continue
            if source_key and trace.progress_report_id and document.get("source_key") == source_key:
                return True
        return False

    @staticmethod
    def _build_pdf_parse_payload(
        extracted_text: str,
        filename: str,
        source_path: str | None,
        extraction: PdfTextExtraction,
        dropbox_metadata: dict | None,
        classification,
    ) -> dict:
        parsed = (
            parse_progress_report(extracted_text).to_public_dict()
            if extracted_text.strip() and classification.status == "accepted"
            else {}
        )
        document = {
            "filename": filename,
            "source_path": source_path,
            "source_key": ImportService._ocr_source_key(filename, source_path),
            "type": classification.document_type,
            "status": classification.status,
            "reason": classification.warnings[0] if classification.warnings else None,
            "score": classification.score,
            "requires_review": classification.requires_review,
            "matched_signals": classification.matched_signals,
            "ignored_signals": classification.ignored_signals,
        }
        if dropbox_metadata:
            document.update(dropbox_metadata)
        parser_logs = parsed.get("logs", [])
        extraction_logs = extraction.page_logs_dict()
        payload = {
            **parsed,
            "document": document,
            "warnings": list(dict.fromkeys(classification.warnings)),
            "page_extraction_log": extraction_logs,
            "extraction_log": [
                *[
                    {
                        "field": "pdf.page",
                        "value": item["chars"],
                        "page": item["page"],
                        "section": None,
                        "method": item["method"],
                        "score": 100 if item["chars"] else 0,
                        "requires_review": item["requires_review"],
                        "warning": item["warning"],
                    }
                    for item in extraction_logs
                ],
                *[
                    {
                        **log,
                        "section": log.get("table"),
                    }
                    for log in parser_logs
                ],
            ],
        }
        return payload

    def _progress_exists_for_source(
        self,
        filename: str,
        source_path: str | None,
        dropbox_metadata: dict | None = None,
    ) -> bool:
        source_key = self._ocr_source_key(filename, source_path)
        dropbox_id = (dropbox_metadata or {}).get("dropbox_id")
        rev = (dropbox_metadata or {}).get("rev")
        traces = (
            self.db.query(ImportedOcrTrace)
            .join(ImportedProgressReport, ImportedOcrTrace.progress_report_id == ImportedProgressReport.id)
            .filter(
                ImportedProgressReport.notes.like(f"{GENERATED_PROGRESS_NOTE}%")
                | ImportedProgressReport.notes.like(f"{LEGACY_GENERATED_PROGRESS_NOTE}%")
            )
            .all()
        )
        if dropbox_id and rev:
            return any(
                (trace.parsed_payload or {}).get("document", {}).get("dropbox_id") == dropbox_id
                and (trace.parsed_payload or {}).get("document", {}).get("rev") == rev
                for trace in traces
            )
        return any(self._ocr_source_key(trace.source_filename, trace.source_path) == source_key for trace in traces)

    @staticmethod
    def _ocr_source_key(filename: str | None, source_path: str | None) -> str:
        value = source_path or filename or ""
        return _normalize_key(value)

    def _build_progress_record_from_pdf_text(
        self,
        job_id: int,
        filename: str,
        source_path: str | None,
        extracted_text: str,
    ) -> ImportedProgressReport | None:
        normalized = _normalize_key(extracted_text)
        if not normalized:
            return None

        careers = self.db.query(Career).all()
        career_name = self._detect_career_name(normalized, careers)
        group_projects = _extract_group_projects(extracted_text)
        group_members = _extract_group_members(extracted_text)
        external_researchers = _extract_external_researchers(extracted_text)
        group_directors = sorted(
            {
                project["director"]
                for project in group_projects
            if project.get("director") and project["director"] != "Director no detectado"
            }
        )
        teacher_name = (
            "Grupo de investigacion"
            if group_projects or group_members or external_researchers
            else self._extract_director_name(extracted_text) or "Docente no detectado"
        )
        year_label = self._detect_year_label(extracted_text, filename, source_path)
        cycle = self._detect_cycle(normalized)
        scientific_products = _extract_all_scientific_products(extracted_text)
        product_counts = _production_counts_from_products(scientific_products)
        articles = product_counts["articles"] if scientific_products else self._detect_articles(normalized)
        books = product_counts["books"] if scientific_products else self._detect_books(normalized)
        book_chapters = (
            product_counts["book_chapters"] if scientific_products else self._detect_book_chapters(normalized)
        )
        presentations = product_counts["presentations"] if scientific_products else self._detect_presentations(normalized)
        projects = 1
        participants = self._extract_participants(extracted_text)
        notes = GENERATED_PROGRESS_NOTE
        if group_projects or group_members or external_researchers:
            member_names = [member["name"] for member in group_members if member.get("name")]
            external_names = [
                f"{researcher['name']} ({researcher['institution']})"
                for researcher in external_researchers
                if researcher.get("name")
            ]
            project_summary = "; ".join(
                f"{project['code']} - {project['name']} ({project['director']})"
                for project in group_projects[:5]
            )
            notes = GENERATED_PROGRESS_NOTE
            if member_names:
                notes += f" Integrantes del GI: {', '.join(member_names[:10])}."
            if external_names:
                notes += f" Investigadores externos: {', '.join(external_names[:5])}."
            if project_summary:
                notes += f" Proyectos detectados del grupo: {project_summary}."
            if group_directors:
                notes += f" Docentes asociados: {', '.join(group_directors)}."
        elif participants:
            notes = f"{GENERATED_PROGRESS_NOTE} Participantes detectados: {', '.join(participants[:8])}."

        return ImportedProgressReport(
            import_job_id=job_id,
            career_name=career_name,
            year_label=year_label,
            cycle=cycle,
            teacher_identifier=None,
            teacher_name=teacher_name,
            articles=articles,
            books=books,
            book_chapters=book_chapters,
            presentations=presentations,
            projects=projects,
            notes=notes,
        )

    def _persist_normalized_from_payload(
        self,
        payload: dict,
        progress: ImportedProgressReport | None,
    ) -> dict[str, int]:
        if not progress:
            return {"teachers": 0, "external_researchers": 0, "research_entities": 0, "projects": 0, "productions": 0}

        if self.db.get_bind().dialect.name == "postgresql":
            self.db.execute(
                text("SELECT pg_advisory_xact_lock(hashtextextended('normalized-import-persistence', 0))")
            )

        period = (
            self.db.query(AcademicPeriod)
            .filter(AcademicPeriod.year_label == progress.year_label, AcademicPeriod.cycle == progress.cycle)
            .first()
        )
        if not period:
            return {"teachers": 0, "external_researchers": 0, "research_entities": 0, "projects": 0, "productions": 0}

        job_id = int(progress.import_job_id) if progress.import_job_id else None
        context = self._normalization_context(progress)
        self._record_parse_review_items(payload, job_id)
        self._record_project_section_gap_audits(payload, job_id, context)
        teachers = self._persist_internal_teachers(
            payload.get("integrantes_internos") or [],
            None,
            job_id,
            context,
            period_id=period.id,
        )
        external_researchers = self._persist_external_researchers(
            payload.get("integrantes_externos") or payload.get("investigadores_externos") or [],
            period.id,
            job_id,
            context,
        )
        self._persist_responsible_roles(payload.get("responsables") or [], progress, period.id, teachers, context)
        research_entities = self._persist_research_entities(
            payload.get("research_entities") or [],
            period.id,
            job_id,
            context,
        )
        primary_research_entity = research_entities[0] if research_entities else None
        projects = self._persist_projects(payload.get("proyectos_fci") or [], period.id, teachers, job_id, context)
        productions = self._persist_productions(
            payload.get("produccion_cientifica") or [],
            period.id,
            teachers,
            job_id=job_id,
            source_section="produccion_cientifica",
            context=context,
            research_entity=primary_research_entity,
        )
        presentations = self._persist_productions(
            payload.get("intercambios") or [],
            period.id,
            teachers,
            default_type=ProductionType.PRESENTATION,
            job_id=job_id,
            source_section="intercambios",
            context=context,
            research_entity=primary_research_entity,
        )
        return {
            "teachers": len(teachers),
            "external_researchers": external_researchers,
            "research_entities": len(research_entities),
            "projects": projects,
            "productions": productions + presentations,
        }

    def _normalization_context(self, progress: ImportedProgressReport | None) -> dict[str, object | None]:
        job = self.db.get(ImportJob, progress.import_job_id) if progress and progress.import_job_id else None
        return {
            "import_batch_id": job.batch_id if job else None,
            "import_job_id": job.id if job else (progress.import_job_id if progress else None),
            "source_file": job.filename if job else None,
            "parser_version": PARSER_VERSION,
        }

    def _record_normalization_audit(
        self,
        context: dict[str, object | None],
        *,
        entity_type: str,
        raw_value: str | None,
        normalized_value: str | None,
        action: str,
        reason: str,
        source_section: str | None,
        confidence_score: float | None,
        normalized_record_id: int | None = None,
        source_page: int | None = None,
        metadata_json: dict | None = None,
    ) -> None:
        import_job_id = context.get("import_job_id")
        if not import_job_id:
            return
        self.db.add(
            ImportNormalizationAudit(
                import_batch_id=context.get("import_batch_id"),
                import_job_id=int(import_job_id),
                normalized_record_id=normalized_record_id,
                entity_type=entity_type,
                source_file=context.get("source_file"),
                source_page=source_page,
                source_section=source_section,
                raw_value=raw_value,
                normalized_value=normalized_value,
                confidence_score=confidence_score,
                parser_version=str(context.get("parser_version") or PARSER_VERSION),
                action=action,
                reason=reason,
                metadata_json=metadata_json,
            )
        )

    def _record_person_role(
        self,
        context: dict[str, object | None],
        *,
        role_type: str,
        person_type: str,
        raw_name: str | None,
        normalized_name: str | None,
        reason: str,
        source_section: str | None,
        confidence_score: float | None,
        validation_status: str | None = None,
        period_id: int | None = None,
        teacher_id: int | None = None,
        external_researcher_id: int | None = None,
        scientific_production_id: int | None = None,
        research_project_id: int | None = None,
        research_entity_id: int | None = None,
        source_page: int | None = None,
        raw_value: str | None = None,
        normalized_value: str | None = None,
        metadata_json: dict | None = None,
    ) -> None:
        import_job_id = context.get("import_job_id")
        if not import_job_id:
            return
        cleaned_name = _clean_person_name(normalized_name or raw_name) or _normalize_text(normalized_name or raw_name)
        person_key = _person_key(cleaned_name)
        resolved_validation_status = validation_status or (
            "validated"
            if person_type in {"teacher", "external_researcher"}
            and bool(teacher_id or external_researcher_id)
            and float(confidence_score or 0) >= 0.90
            else "pending_review"
        )
        exists = (
            self.db.query(PersonRole)
            .filter(
                PersonRole.import_job_id == int(import_job_id),
                PersonRole.role_type == role_type,
                PersonRole.person_key == person_key,
                PersonRole.source_section == source_section,
                PersonRole.scientific_production_id == scientific_production_id,
                PersonRole.research_project_id == research_project_id,
                PersonRole.research_entity_id == research_entity_id,
            )
            .first()
        )
        if exists:
            if resolved_validation_status == "validated" and exists.validation_status != "validated":
                exists.validation_status = "validated"
                exists.person_type = person_type
                exists.teacher_id = teacher_id or exists.teacher_id
                exists.external_researcher_id = external_researcher_id or exists.external_researcher_id
                exists.normalized_name = cleaned_name or exists.normalized_name
                exists.confidence_score = max(float(exists.confidence_score or 0), float(confidence_score or 0))
                exists.reason = reason
                exists.metadata_json = metadata_json or exists.metadata_json
            return
        self.db.add(
            PersonRole(
                period_id=period_id,
                import_batch_id=context.get("import_batch_id"),
                import_job_id=int(import_job_id),
                teacher_id=teacher_id,
                external_researcher_id=external_researcher_id,
                scientific_production_id=scientific_production_id,
                research_project_id=research_project_id,
                research_entity_id=research_entity_id,
                role_type=role_type,
                person_type=person_type,
                person_key=person_key or None,
                raw_name=raw_name,
                normalized_name=cleaned_name or normalized_name,
                source_file=context.get("source_file"),
                source_page=source_page,
                source_section=source_section,
                raw_value=raw_value or raw_name,
                normalized_value=normalized_value or cleaned_name or normalized_name,
                confidence_score=confidence_score,
                reason=reason,
                parser_version=str(context.get("parser_version") or PARSER_VERSION),
                metadata_json=metadata_json,
                validation_status=resolved_validation_status,
            )
        )

    def _record_parse_review_items(self, payload: dict, job_id: int | None) -> None:
        if not job_id:
            return
        for log in payload.get("logs") or []:
            if not isinstance(log, dict) or not log.get("requires_review"):
                continue
            field = _normalize_text(log.get("field")) or "parser"
            raw_value = _normalize_text(log.get("value"))
            source_section = _normalize_text(log.get("table") or log.get("section"))
            source_page = _optional_int(log.get("page"))
            score = log.get("score")
            confidence_score = None
            if isinstance(score, (int, float)):
                confidence_score = float(score) / 100 if score > 1 else float(score)
            self._record_review_item(
                job_id,
                field=field,
                raw_value=raw_value,
                normalized_value=None,
                reason="Seccion o campo detectado con baja confianza; requiere revision de contexto.",
                source_section=source_section,
                confidence_score=confidence_score,
                source_page=source_page,
            )

    def _record_project_section_gap_audits(
        self,
        payload: dict,
        job_id: int | None,
        context: dict[str, object | None],
    ) -> None:
        if not job_id or payload.get("proyectos_fci"):
            return
        seen: set[tuple[str, int | None]] = set()
        for log in payload.get("logs") or []:
            if not isinstance(log, dict):
                continue
            source_section = _normalize_text(log.get("table") or log.get("section"))
            if source_section != "proyectos_fci":
                continue
            raw_value = _normalize_text(log.get("value"))
            source_page = _optional_int(log.get("page"))
            key = (raw_value, source_page)
            if key in seen:
                continue
            seen.add(key)
            reason = (
                "Se detecto seccion oficial de proyectos FCI, pero no se encontraron filas "
                "con codigo/nombre/director/avance/estado suficientes."
            )
            metadata = {
                "section_detected": raw_value or None,
                "candidates_found": 0,
                "persisted": 0,
                "requires_review": 1,
                "discarded": 0,
                "reason": reason,
            }
            self._record_normalization_audit(
                context,
                entity_type="research_project",
                raw_value=raw_value,
                normalized_value=None,
                action="requires_review",
                reason=reason,
                source_section="proyectos_fci",
                confidence_score=0.5,
                source_page=source_page,
                metadata_json=metadata,
            )
            self._record_review_item(
                job_id,
                field="research_project",
                raw_value=raw_value,
                normalized_value=None,
                reason=reason,
                source_section="proyectos_fci",
                confidence_score=0.5,
                source_page=source_page,
            )

    def _accepted_pdf_review_state(
        self,
        classification_score: float,
        parsed_payload: dict,
        progress: ImportedProgressReport,
    ) -> tuple[str, str | None]:
        issues: list[str] = []
        if classification_score < 75:
            issues.append("Confianza de clasificacion baja.")
        if not progress.career_name:
            issues.append("No se detecto carrera del informe.")
        if progress.teacher_name in {"Docente no detectado", ""}:
            issues.append("No se detecto docente o grupo responsable.")
        has_structured_data = any(
            parsed_payload.get(key)
            for key in (
                "integrantes_internos",
                "integrantes_externos",
                "proyectos_fci",
                "research_entities",
                "produccion_cientifica",
                "intercambios",
            )
        )
        if not has_structured_data:
            issues.append("No se detectaron secciones estructuradas del informe.")
        if issues:
            return "PENDIENTE_REVISION", " ".join(issues)
        return "VALIDADO_AUTOMATICO", None

    def _persist_internal_teachers(
        self,
        members: list[dict],
        fallback_career_name: str | None,
        job_id: int | None = None,
        context: dict[str, object | None] | None = None,
        period_id: int | None = None,
    ) -> list[Teacher]:
        persisted: list[Teacher] = []
        context = context or {}
        for member in members:
            raw_name = _normalize_text(member.get("name"))
            source_section = _normalize_text(member.get("source_section")) or "integrantes_internos"
            source_page = _optional_int(member.get("source_page"))
            name = _clean_person_name(raw_name)
            raw_faculty = _normalize_text(member.get("faculty"))
            normalized_faculty = _normalize_text(raw_faculty)
            raw_career = _normalize_text(member.get("career")) or fallback_career_name
            normalized_name = _person_key(name or raw_name)
            teacher_metadata = {
                "raw_name": raw_name or None,
                "normalized_name": name or None,
                "raw_faculty": raw_faculty or None,
                "normalized_faculty": normalized_faculty or None,
                "raw_career": raw_career or None,
                "normalized_career": None,
            }
            if not name:
                self._record_normalization_audit(
                    context,
                    entity_type="teacher",
                    raw_value=raw_name,
                    normalized_value=None,
                    action="discarded_invalid",
                    reason="Nombre de docente interno ausente o no confiable.",
                    source_section=source_section,
                    confidence_score=0.4,
                    source_page=source_page,
                    metadata_json=teacher_metadata,
                )
                self._record_review_item(
                    job_id,
                    field="teacher",
                    raw_value=raw_name,
                    normalized_value=None,
                    reason="Docente interno descartado por nombre ausente o no confiable.",
                    source_section=source_section,
                    confidence_score=0.4,
                    source_page=source_page,
                )
                continue
            career_name = raw_career
            career = self._find_career(career_name)
            if not career:
                teacher_metadata["person_type"] = "unresolved"
                teacher_metadata["dedupe_signature"] = "|".join([normalized_name, _normalize_key(raw_career)])
                teacher_metadata["dedupe_strategy"] = "normalized_name+raw_career_requires_catalog_review"
                self._record_normalization_audit(
                    context,
                    entity_type="teacher",
                    raw_value=raw_name,
                    normalized_value=name,
                    action="requires_review",
                    reason="Carrera interna no identificada desde la seccion oficial.",
                    source_section=source_section,
                    confidence_score=0.5,
                    source_page=source_page,
                    metadata_json=teacher_metadata,
                )
                self._record_person_role(
                    context,
                    role_type="integrante_interno",
                    person_type="unresolved",
                    raw_name=raw_name,
                    normalized_name=name,
                    reason="Integrante interno detectado, pero la carrera no existe en el catalogo actual.",
                    source_section=source_section,
                    confidence_score=0.5,
                    period_id=period_id,
                    source_page=source_page,
                    raw_value=raw_name,
                    normalized_value=name,
                    metadata_json=teacher_metadata,
                )
                self._record_review_item(
                    job_id,
                    field="career",
                    raw_value=career_name,
                    normalized_value=None,
                    reason="Carrera interna no identificada contra catalogo.",
                    source_section=source_section,
                    confidence_score=0.0,
                    source_page=source_page,
                )
                continue
            normalized_career = career.name
            teacher_metadata["normalized_career"] = normalized_career
            dedupe_signature = "|".join([normalized_name, _normalize_key(normalized_career)])
            teacher_metadata["dedupe_signature"] = dedupe_signature
            teacher_metadata["dedupe_strategy"] = "normalized_name+normalized_career"
            email = self._synthetic_teacher_email(name, career.id)
            teacher = self.db.query(Teacher).filter(Teacher.institutional_email == email).first()
            if not teacher:
                teacher = Teacher(
                    career_id=career.id,
                    full_name=name,
                    institutional_email=email,
                    research_hours=0,
                    is_active=True,
                    import_batch_id=context.get("import_batch_id"),
                    import_job_id=context.get("import_job_id"),
                    source_file=context.get("source_file"),
                    source_page=source_page,
                    source_section=source_section,
                    raw_value=raw_name,
                    normalized_value=name,
                    confidence_score=0.95,
                    parser_version=PARSER_VERSION,
                    normalization_action="persisted",
                    normalization_reason="Docente interno detectado en seccion oficial y carrera validada contra catalogo.",
                    raw_name=raw_name,
                    normalized_name=normalized_name,
                    raw_faculty=raw_faculty,
                    normalized_faculty=normalized_faculty,
                    raw_career=raw_career,
                    normalized_career=normalized_career,
                )
                self.db.add(teacher)
                self.db.flush()
                teacher_metadata["matched_teacher_id"] = teacher.id
                self._record_normalization_audit(
                    context,
                    entity_type="teacher",
                    raw_value=raw_name,
                    normalized_value=name,
                    action="persisted",
                    reason="Docente interno detectado en seccion oficial y carrera validada contra catalogo.",
                    source_section=source_section,
                    confidence_score=0.95,
                    normalized_record_id=teacher.id,
                    source_page=source_page,
                    metadata_json=teacher_metadata,
                )
            else:
                teacher_metadata["matched_teacher_id"] = teacher.id
                self._record_normalization_audit(
                    context,
                    entity_type="teacher",
                    raw_value=raw_name,
                    normalized_value=name,
                    action="merged_duplicate",
                    reason="Docente ya existia con la misma firma normalizada nombre+carrera.",
                    source_section=source_section,
                    confidence_score=0.95,
                    normalized_record_id=teacher.id,
                    source_page=source_page,
                    metadata_json=teacher_metadata,
                )
            self._record_person_role(
                context,
                role_type="integrante_interno",
                person_type="teacher",
                raw_name=raw_name,
                normalized_name=name,
                reason="Rol detectado desde seccion oficial de integrantes internos.",
                source_section=source_section,
                confidence_score=0.95,
                period_id=period_id,
                teacher_id=teacher.id,
                source_page=source_page,
                raw_value=raw_name,
                normalized_value=name,
                metadata_json=teacher_metadata,
            )
            persisted.append(teacher)
        return persisted

    def _persist_external_researchers(
        self,
        researchers: list[dict],
        period_id: int,
        job_id: int | None,
        context: dict[str, object | None] | None = None,
    ) -> int:
        created_or_existing = 0
        context = context or {}
        for researcher in researchers:
            raw_name = _normalize_text(researcher.get("name"))
            name = _clean_person_name(raw_name)
            raw_institution = _normalize_text(researcher.get("institution"))
            source_section = _normalize_text(researcher.get("source_section")) or "integrantes_externos"
            source_page = _optional_int(researcher.get("source_page"))
            raw_participant_type = _normalize_text(researcher.get("participant_type"))
            participant_type = _external_participant_type(
                " ".join(
                    filter(
                        None,
                        [
                            raw_name,
                            raw_institution,
                            raw_participant_type,
                            _normalize_text(researcher.get("reason")),
                        ],
                    )
                ),
                default=raw_participant_type or "investigador_externo",
            )
            institution = _clean_external_institution(raw_institution)
            assigned_person_type = "investigador_externo" if institution else participant_type
            raw_value = f"{raw_name or ''} | {raw_institution or ''}".strip()
            external_metadata = {
                "raw_name": raw_name or None,
                "normalized_name": name or None,
                "raw_institution": raw_institution or None,
                "normalized_institution": institution or None,
                "raw_participant_type": raw_participant_type or None,
                "participant_type": assigned_person_type,
            }
            if not name or not institution:
                action = "requires_review" if name and participant_type in {"estudiante", "graduado"} else "discarded_invalid"
                reason = (
                    "Participante estudiante/graduado detectado; se conserva como rol en revision, no como docente interno."
                    if action == "requires_review"
                    else "Investigador externo descartado por nombre/institucion dudosa o institucion no externa."
                )
                self._record_normalization_audit(
                    context,
                    entity_type="external_researcher",
                    raw_value=raw_value,
                    normalized_value=None,
                    action=action,
                    reason=reason,
                    source_section=source_section,
                    confidence_score=0.5,
                    source_page=source_page,
                    metadata_json=external_metadata,
                )
                if name and participant_type in {"estudiante", "graduado"}:
                    self._record_person_role(
                        context,
                        role_type="participante_externo",
                        person_type=participant_type,
                        raw_name=raw_name,
                        normalized_name=name,
                        reason=reason,
                        source_section=source_section,
                        confidence_score=0.65,
                        period_id=period_id,
                        source_page=source_page,
                        raw_value=raw_value,
                        normalized_value=name,
                        metadata_json=external_metadata,
                    )
                self._record_review_item(
                    job_id,
                    field="external_researcher",
                    raw_value=raw_value,
                    normalized_value=None,
                    reason=reason,
                    source_section=source_section,
                    confidence_score=0.5,
                    source_page=source_page,
                )
                continue
            normalized_name = _person_key(name)
            normalized_institution = _normalize_key(institution)
            normalized_value = f"{name} | {institution}"
            external_metadata["dedupe_signature"] = f"{normalized_name}|{normalized_institution}|{period_id}"
            external_metadata["dedupe_strategy"] = "normalized_name+normalized_institution+period"
            researcher_row = (
                self.db.query(ExternalResearcher)
                .filter(
                    ExternalResearcher.period_id == period_id,
                    ExternalResearcher.normalized_name == normalized_name,
                    ExternalResearcher.normalized_institution == normalized_institution,
                )
                .first()
            )
            if not researcher_row:
                researcher_row = ExternalResearcher(
                    period_id=period_id,
                    import_batch_id=context.get("import_batch_id"),
                    import_job_id=job_id,
                    source_file=context.get("source_file"),
                    source_page=source_page,
                    full_name=name,
                    normalized_name=normalized_name,
                    institution=institution,
                    normalized_institution=normalized_institution,
                    source_section=source_section,
                    confidence_score=0.9,
                    requires_review=False,
                    raw_value=raw_value,
                    normalized_value=normalized_value,
                    parser_version=PARSER_VERSION,
                    normalization_action="persisted",
                    normalization_reason="Investigador externo detectado en seccion oficial y universidad externa validada.",
                    raw_name=raw_name,
                    raw_institution=raw_institution,
                )
                self.db.add(researcher_row)
                self.db.flush()
                action = "persisted"
                reason = "Investigador externo detectado en seccion oficial y universidad externa validada."
            else:
                action = "merged_duplicate"
                reason = "Investigador externo ya existia con la misma firma nombre+institucion+periodo."
            external_metadata["matched_external_researcher_id"] = researcher_row.id
            self._record_normalization_audit(
                context,
                entity_type="external_researcher",
                raw_value=raw_value,
                normalized_value=normalized_value,
                action=action,
                reason=reason,
                source_section=source_section,
                confidence_score=0.9,
                normalized_record_id=researcher_row.id,
                source_page=source_page,
                metadata_json=external_metadata,
            )
            self._record_person_role(
                context,
                role_type="participante_externo",
                person_type=assigned_person_type,
                raw_name=raw_name,
                normalized_name=name,
                reason="Rol externo detectado desde seccion oficial de investigadores/participantes externos.",
                source_section=source_section,
                confidence_score=0.9,
                period_id=period_id,
                external_researcher_id=researcher_row.id,
                source_page=source_page,
                raw_value=raw_value,
                normalized_value=normalized_value,
                metadata_json=external_metadata,
            )
            created_or_existing += 1
        return created_or_existing

    def _record_review_item(
        self,
        job_id: int | None,
        *,
        field: str,
        raw_value: str | None,
        normalized_value: str | None,
        reason: str,
        source_section: str | None = None,
        confidence_score: float | None = None,
        source_page: int | None = None,
    ) -> None:
        if not job_id:
            return
        exists = (
            self.db.query(ImportReviewItem)
            .filter(
                ImportReviewItem.import_job_id == job_id,
                ImportReviewItem.field == field,
                ImportReviewItem.raw_value == raw_value,
                ImportReviewItem.reason == reason,
            )
            .first()
        )
        if exists:
            return
        self.db.add(
            ImportReviewItem(
                import_job_id=job_id,
                field=field,
                source_page=source_page,
                source_section=source_section,
                raw_value=raw_value,
                normalized_value=normalized_value,
                confidence_score=confidence_score,
                reason=reason,
            )
        )

    def _persist_responsible_roles(
        self,
        responsible_items: list[dict],
        progress: ImportedProgressReport,
        period_id: int,
        teachers: list[Teacher],
        context: dict[str, object | None],
    ) -> None:
        items = list(responsible_items)
        fallback_name = _normalize_text(progress.teacher_name)
        if fallback_name and fallback_name not in {"Docente no detectado", "Director no detectado"}:
            items.append(
                {
                    "name": fallback_name,
                    "role_type": "responsable_informe",
                    "source_section": "imported_progress_report",
                    "source_page": None,
                    "reason": "Responsable inferido desde el resumen persistido del informe.",
                }
            )

        external_researchers = (
            self.db.query(ExternalResearcher)
            .filter(ExternalResearcher.period_id == period_id, ExternalResearcher.requires_review.is_(False))
            .all()
        )
        seen: set[tuple[str, str, str | None]] = set()
        for item in items:
            raw_name = _normalize_text(item.get("name"))
            name = _clean_person_name(raw_name)
            role_type = _normalize_text(item.get("role_type")) or "responsable_informe"
            source_section = _normalize_text(item.get("source_section")) or "director_responsable"
            source_page = _optional_int(item.get("source_page"))
            key = (role_type, _person_key(name or raw_name), source_section)
            if key in seen:
                continue
            seen.add(key)
            if not name:
                continue
            tokens = _person_match_tokens(name)
            teacher, teacher_confidence = (
                self._match_teacher_by_tokens(teachers, tokens, name) if len(tokens) >= 2 else (None, 0.0)
            )
            external, external_confidence = (
                self._match_external_by_tokens(external_researchers, tokens, name) if len(tokens) >= 2 else (None, 0.0)
            )
            person_type = "teacher" if teacher else ("external_researcher" if external else "unresolved")
            confidence = teacher_confidence or external_confidence or 0.55
            action = "persisted" if teacher or external else "requires_review"
            reason = (
                "Rol vinculado a docente interno detectado en el mismo informe."
                if teacher
                else (
                    "Rol vinculado a investigador externo detectado para el periodo."
                    if external
                    else "Rol detectado, pero la persona no coincide con docente interno ni investigador externo."
                )
            )
            metadata = {
                "role_type": role_type,
                "raw_name": raw_name,
                "normalized_name": name,
                "person_type": person_type,
                "matched_teacher_id": teacher.id if teacher else None,
                "matched_external_researcher_id": external.id if external else None,
                "source": item.get("reason") or "section_context",
            }
            self._record_person_role(
                context,
                role_type=role_type,
                person_type=person_type,
                raw_name=raw_name,
                normalized_name=name,
                reason=reason,
                source_section=source_section,
                confidence_score=confidence,
                period_id=period_id,
                teacher_id=teacher.id if teacher else None,
                external_researcher_id=external.id if external else None,
                source_page=source_page,
                raw_value=raw_name,
                normalized_value=name,
                metadata_json=metadata,
            )
            self._record_normalization_audit(
                context,
                entity_type="person_role",
                raw_value=raw_name,
                normalized_value=name,
                action=action,
                reason=reason,
                source_section=source_section,
                confidence_score=confidence,
                normalized_record_id=teacher.id if teacher else (external.id if external else None),
                source_page=source_page,
                metadata_json=metadata,
            )
            if not teacher and not external:
                self._record_review_item(
                    int(context.get("import_job_id")) if context.get("import_job_id") else None,
                    field="person_role",
                    raw_value=raw_name,
                    normalized_value=name,
                    reason=reason,
                    source_section=source_section,
                    confidence_score=confidence,
                    source_page=source_page,
                )

    def _persist_research_entities(
        self,
        entities: list[dict],
        period_id: int,
        job_id: int | None,
        context: dict[str, object | None] | None = None,
    ) -> list[ResearchEntity]:
        persisted_entities: list[ResearchEntity] = []
        context = context or {}
        for item in entities:
            if not isinstance(item, dict):
                continue
            entity_type = _normalize_text(item.get("type")) or "pendiente_clasificacion"
            raw_code = _normalize_text(item.get("code"))[:120]
            normalized_code = (_normalize_key(raw_code).replace(" ", "") or None)
            normalized_code = normalized_code[:140] if normalized_code else None
            raw_name = _normalize_text(item.get("title") or item.get("name"))[:350]
            normalized_name = _normalize_key(raw_name) or None
            normalized_name = normalized_name[:380] if normalized_name else None
            raw_director = _normalize_text(item.get("director"))[:220]
            normalized_director = _clean_person_name(raw_director) or _normalize_text(raw_director) or None
            normalized_director = normalized_director[:260] if normalized_director else None
            raw_status = _normalize_text(item.get("status"))[:80]
            source_section = _normalize_text(item.get("source_section")) or "datos_generales"
            source_page = _optional_int(item.get("source_page"))
            confidence = item.get("confidence_score")
            confidence_score = float(confidence) if isinstance(confidence, (int, float)) else 0.5
            reason = _normalize_text(item.get("reason")) or "Entidad investigativa detectada desde el PDF."
            metadata = {
                "raw_type": entity_type,
                "raw_code": raw_code or None,
                "normalized_code": normalized_code,
                "raw_name": raw_name or None,
                "normalized_name": normalized_name,
                "raw_director": raw_director or None,
                "normalized_director": normalized_director,
                "raw_status": raw_status or None,
                "normalized_status": raw_status or None,
                "raw_career": (_normalize_text(item.get("career")) or None),
                "normalized_career": (_clean_career_name(item.get("career")) or _normalize_text(item.get("career")) or None),
                "source": "research_entities",
            }
            raw_value = json.dumps({**item, "normalized_code": normalized_code, "normalized_name": normalized_name}, ensure_ascii=False)

            if entity_type == "documento_no_relevante" or item.get("validation_status") == "discarded_invalid":
                self._record_normalization_audit(
                    context,
                    entity_type="research_entity",
                    raw_value=raw_value,
                    normalized_value=None,
                    action="discarded_invalid",
                    reason=reason,
                    source_section=source_section,
                    confidence_score=confidence_score,
                    source_page=source_page,
                    metadata_json=metadata,
                )
                continue

            if not normalized_code and not normalized_name:
                self._record_normalization_audit(
                    context,
                    entity_type="research_entity",
                    raw_value=raw_value,
                    normalized_value=None,
                    action="pending_review",
                    reason="Entidad investigativa detectada sin codigo ni nombre suficientes.",
                    source_section=source_section,
                    confidence_score=confidence_score,
                    source_page=source_page,
                    metadata_json=metadata,
                )
                self._record_review_item(
                    job_id,
                    field="research_entity",
                    raw_value=raw_value,
                    normalized_value=None,
                    reason="Entidad investigativa detectada sin codigo ni nombre suficientes.",
                    source_section=source_section,
                    confidence_score=confidence_score,
                    source_page=source_page,
                )
                continue

            query = self.db.query(ResearchEntity).filter(
                ResearchEntity.period_id == period_id,
                ResearchEntity.type == entity_type,
            )
            if normalized_code:
                query = query.filter(ResearchEntity.normalized_code == normalized_code)
            else:
                query = query.filter(ResearchEntity.normalized_name == normalized_name)
            entity = query.first()
            validation_status = (
                "validated"
                if item.get("validation_status") == "validated" and (normalized_code or normalized_name) and normalized_director
                else "pending_review"
            )
            if not entity:
                entity = ResearchEntity(
                    period_id=period_id,
                    type=entity_type,
                    code=raw_code or None,
                    normalized_code=normalized_code,
                    name=raw_name or None,
                    normalized_name=normalized_name,
                    director_name=raw_director or None,
                    normalized_director_name=normalized_director,
                    year=_optional_int(item.get("year")),
                    cycle=_optional_int(item.get("cycle")),
                    academic_unit=(_normalize_text(item.get("academic_unit")) or None)[:220] if _normalize_text(item.get("academic_unit")) else None,
                    career_name=(_clean_career_name(item.get("career")) or _normalize_text(item.get("career")) or None)[:220] if (_clean_career_name(item.get("career")) or _normalize_text(item.get("career"))) else None,
                    progress_percentage=self._progress_percent(item.get("progress_percentage")),
                    status=raw_status or None,
                    validation_status=validation_status,
                    import_batch_id=context.get("import_batch_id"),
                    import_job_id=job_id or context.get("import_job_id"),
                    source_file=context.get("source_file"),
                    source_page=source_page,
                    source_section=source_section,
                    raw_value=raw_value,
                    normalized_value=normalized_code or normalized_name,
                    confidence_score=confidence_score,
                    parser_version=str(context.get("parser_version") or PARSER_VERSION),
                    reason=reason,
                    metadata_json=metadata,
                )
                self.db.add(entity)
                self.db.flush()
                action = "persisted" if validation_status == "validated" else "pending_review"
                audit_reason = reason if validation_status == "validated" else f"{reason} Requiere validacion academica."
            else:
                action = "merged_duplicate"
                audit_reason = "Entidad investigativa fusionada por tipo y codigo/nombre normalizado en el mismo periodo."
                if validation_status == "validated" and entity.validation_status != "validated":
                    entity.validation_status = "validated"
                    entity.director_name = raw_director or entity.director_name
                    entity.normalized_director_name = normalized_director or entity.normalized_director_name
                    entity.reason = reason
                if not entity.import_batch_id and context.get("import_batch_id"):
                    entity.import_batch_id = context.get("import_batch_id")
                if not entity.import_job_id and (job_id or context.get("import_job_id")):
                    entity.import_job_id = job_id or context.get("import_job_id")
                if not entity.source_file and context.get("source_file"):
                    entity.source_file = context.get("source_file")

            self._record_normalization_audit(
                context,
                entity_type="research_entity",
                raw_value=raw_value,
                normalized_value=normalized_code or normalized_name,
                action=action,
                reason=audit_reason,
                source_section=source_section,
                confidence_score=confidence_score,
                normalized_record_id=entity.id,
                source_page=source_page,
                metadata_json=metadata,
            )
            if validation_status == "pending_review":
                self._record_review_item(
                    job_id,
                    field="research_entity",
                    raw_value=raw_value,
                    normalized_value=normalized_code or normalized_name,
                    reason=audit_reason,
                    source_section=source_section,
                    confidence_score=confidence_score,
                    source_page=source_page,
                )
            if normalized_director:
                role_type = {
                    "semillero": "tutor_semillero",
                    "grupo_investigacion": "coordinador",
                    "proyecto_fci": "director",
                    "informe_seguimiento": "responsable_informe",
                }.get(entity_type, "responsable_informe")
                self._record_person_role(
                    context,
                    role_type=role_type,
                    person_type="unresolved",
                    raw_name=raw_director,
                    normalized_name=normalized_director,
                    reason="Rol detectado desde datos generales de la entidad investigativa.",
                    source_section=source_section,
                    confidence_score=confidence_score,
                    period_id=period_id,
                    research_entity_id=entity.id,
                    source_page=source_page,
                    raw_value=raw_director,
                    normalized_value=normalized_director,
                    metadata_json={**metadata, "research_entity_id": entity.id},
                )
            persisted_entities.append(entity)
        return persisted_entities

    def _persist_projects(
        self,
        projects: list[dict],
        period_id: int,
        teachers: list[Teacher],
        job_id: int | None = None,
        context: dict[str, object | None] | None = None,
    ) -> int:
        persisted = 0
        context = context or {}
        for item in projects:
            name = _normalize_text(item.get("name")) or _normalize_text(item.get("title"))
            raw_code = _normalize_text(item.get("code"))
            raw_status = _normalize_text(item.get("status"))
            raw_progress = _normalize_text(item.get("progress"))
            source_section = _normalize_text(item.get("source_section")) or "proyectos_fci"
            source_page = _optional_int(item.get("source_page"))
            if not name:
                self._record_normalization_audit(
                    context,
                    entity_type="research_project",
                    raw_value=str(item),
                    normalized_value=None,
                    action="discarded_invalid",
                    reason="Proyecto descartado por falta de nombre/titulo.",
                    source_section=source_section,
                    confidence_score=0.4,
                    source_page=source_page,
                )
                continue
            normalized_project_name = _normalize_key(name)
            normalized_code = _normalize_key(raw_code)
            normalized_status = _normalize_text(raw_status) or "vigente"
            normalized_progress = self._progress_percent(raw_progress) or 0
            project_metadata = {
                "raw_project_name": name or None,
                "normalized_project_name": normalized_project_name or None,
                "raw_code": raw_code or None,
                "normalized_code": normalized_code or None,
                "raw_status": raw_status or None,
                "normalized_status": normalized_status or None,
                "raw_progress": raw_progress or None,
                "normalized_progress": normalized_progress,
            }
            project = (
                self.db.query(ResearchProject)
                .filter(
                    ResearchProject.period_id == period_id,
                    ResearchProject.normalized_project_name == normalized_project_name,
                )
                .first()
            )
            if not project:
                project = ResearchProject(
                    period_id=period_id,
                    name=name[:250],
                    project_type=ProjectType.FCI,
                    description=raw_code,
                    status=normalized_status,
                    progress_percentage=normalized_progress,
                    import_batch_id=context.get("import_batch_id"),
                    import_job_id=job_id,
                    source_file=context.get("source_file"),
                    source_page=source_page,
                    source_section=source_section,
                    raw_value=name,
                    normalized_value=name,
                    confidence_score=0.9,
                    parser_version=PARSER_VERSION,
                    normalization_action="persisted",
                    normalization_reason="Proyecto detectado en seccion oficial de proyectos FCI.",
                    raw_project_name=name,
                    normalized_project_name=normalized_project_name,
                    raw_code=raw_code,
                    normalized_code=normalized_code,
                    raw_status=raw_status,
                    normalized_status=normalized_status,
                    raw_progress=raw_progress,
                    normalized_progress=normalized_progress,
                )
                self.db.add(project)
                self.db.flush()
                action = "persisted"
                reason = "Proyecto detectado en seccion oficial de proyectos FCI."
            else:
                action = "merged_duplicate"
                reason = "Proyecto ya existia con el mismo nombre normalizado y periodo."
            self._record_normalization_audit(
                context,
                entity_type="research_project",
                raw_value=name,
                normalized_value=name,
                action=action,
                reason=reason,
                source_section=source_section,
                confidence_score=0.9,
                normalized_record_id=project.id,
                source_page=source_page,
                metadata_json=project_metadata,
            )
            persisted += 1
            director_name = _clean_person_name(_normalize_text(item.get("director")))
            if director_name:
                director = self._match_teacher_by_tokens(teachers, _person_match_tokens(director_name), director_name)
                role_metadata = {
                    "role_type": "director",
                    "raw_name": _normalize_text(item.get("director")),
                    "normalized_name": director_name,
                    "matched_teacher_id": director.id if director else None,
                    "project_id": project.id,
                    "project_name": project.name,
                }
                self._record_person_role(
                    context,
                    role_type="director",
                    person_type="teacher" if director else "unresolved",
                    raw_name=_normalize_text(item.get("director")),
                    normalized_name=director_name,
                    reason=(
                        "Director detectado explicitamente en la fila/bloque del proyecto FCI."
                        if director
                        else "Director detectado en proyecto FCI, pero no coincide con integrante interno del informe."
                    ),
                    source_section=source_section,
                    confidence_score=0.9 if director else 0.55,
                    period_id=period_id,
                    teacher_id=director.id if director else None,
                    research_project_id=project.id,
                    source_page=source_page,
                    raw_value=_normalize_text(item.get("director")),
                    normalized_value=director_name,
                    metadata_json=role_metadata,
                )
                if director:
                    exists = (
                        self.db.query(ProjectTeacher)
                        .filter(ProjectTeacher.project_id == project.id, ProjectTeacher.teacher_id == director.id)
                        .first()
                    )
                    if not exists:
                        self.db.add(
                            ProjectTeacher(
                                project_id=project.id,
                                teacher_id=director.id,
                                role=ProjectTeacherRole.DIRECTOR,
                            )
                        )
                else:
                    self._record_review_item(
                        job_id,
                        field="project_director",
                        raw_value=_normalize_text(item.get("director")),
                        normalized_value=director_name,
                        reason="Director de proyecto FCI no coincide con integrante interno detectado.",
                        source_section=source_section,
                        confidence_score=0.55,
                        source_page=source_page,
                    )
        return persisted

    def _persist_productions(
        self,
        products: list[dict],
        period_id: int,
        teachers: list[Teacher],
        default_type: ProductionType | None = None,
        job_id: int | None = None,
        source_section: str = "produccion_cientifica",
        context: dict | None = None,
        research_entity: ResearchEntity | None = None,
    ) -> int:
        persisted = 0
        for item in products:
            raw_title = _normalize_text(item.get("title"))
            authors = item.get("authors") or []
            title = clean_product_title(raw_title, authors) or raw_title
            title_was_cleaned = bool(raw_title and title and _normalize_key(raw_title) != _normalize_key(title))
            item_source_section = _normalize_text(item.get("source_section")) or source_section
            item_source_field = _normalize_text(item.get("source_field")) or "title"
            row_or_block_id = _normalize_text(item.get("row_or_block_id"))
            source_page = _optional_int(item.get("source_page"))
            raw_authors = json.dumps(authors, ensure_ascii=False) if authors else None
            normalized_authors = ", ".join(
                filter(None, (_clean_person_name(str(author)) or _normalize_text(author) for author in authors))
            )
            raw_status = _normalize_text(item.get("status"))
            normalized_status = self._production_status_from_payload(item.get("status"))
            raw_impact = _normalize_text(item.get("impact"))
            normalized_impact = _normalize_text(item.get("impact"))
            normalized_title = _normalize_key(title)
            resolved_authors = self._resolve_product_authors(
                authors,
                teachers,
                period_id,
                context=context,
                source_section=item_source_section,
                source_field="authors",
                row_or_block_id=row_or_block_id,
                authors_source=_normalize_text(item.get("authors_source")) or None,
            )
            production_metadata = {
                "raw_title": raw_title or None,
                "raw_text": item.get("original_text") or item.get("raw_title") or raw_title or None,
                "clean_title": title or None,
                "normalized_title": normalized_title or None,
                "source_field": item_source_field,
                "row_or_block_id": row_or_block_id or None,
                "fields_detected": item.get("fields_detected") or [],
                "fields_missing": item.get("fields_missing") or [],
                "raw_status": raw_status or None,
                "normalized_status": normalized_status or None,
                "raw_impact": raw_impact or None,
                "normalized_impact": normalized_impact or None,
                "authors": resolved_authors,
            }
            title_quality, title_quality_reason = product_title_quality(title)
            production_metadata["title_quality"] = title_quality
            if title_quality == "invalid":
                self._record_normalization_audit(
                    context,
                    entity_type="scientific_production",
                    raw_value=title or str(item),
                    normalized_value=None,
                    action="discarded_invalid",
                    reason=title_quality_reason,
                    source_section=item_source_section,
                    confidence_score=0.25,
                    source_page=source_page,
                    metadata_json=production_metadata,
                )
                self._record_review_item(
                    job_id,
                    field="scientific_production",
                    raw_value=title or str(item),
                    normalized_value=None,
                    reason=title_quality_reason,
                    source_section=item_source_section,
                    confidence_score=0.25,
                    source_page=source_page,
                )
                continue
            teacher = self._primary_internal_author(resolved_authors, teachers)
            review_reasons: list[str] = []
            title_requires_review = title_quality == "truncated" or item.get("requires_review")
            if title_quality == "truncated":
                review_reasons.append(title_quality_reason)
            if title_was_cleaned:
                title_requires_review = True
                review_reasons.append("titulo mezclado con autores/OCR; se conserva texto crudo para revision.")
            if item.get("requires_review"):
                if title_quality_reason not in review_reasons:
                    review_reasons.append("titulo truncado o incompleto")
            if _normalize_key(item.get("validation_status")) in {"PENDING_REVIEW", "PENDIENTE_VALIDACION"}:
                if "Producto detectado pendiente de validacion academica." not in review_reasons:
                    review_reasons.append("Producto detectado pendiente de validacion academica.")
            if not teachers:
                review_reasons.append("Producto detectado sin docentes internos confiables en el informe.")
            if not teacher:
                review_reasons.append("Producto detectado sin autor interno resuelto.")
            if title_requires_review and normalized_status == "published":
                normalized_status = "pending_review"
            validation_status = "pending_review" if review_reasons else "validated"
            confidence_score = 0.55 if review_reasons else 0.9
            review_reason = " ".join(review_reasons) or None
            production_type = default_type or self._production_type_from_payload(item.get("type"))
            production = (
                self.db.query(ScientificProduction)
                .filter(
                    ScientificProduction.period_id == period_id,
                    ScientificProduction.normalized_title == normalized_title[:300],
                )
                .first()
            )
            if not production:
                production = ScientificProduction(
                    teacher_id=teacher.id if teacher else None,
                    period_id=period_id,
                    research_entity_id=research_entity.id if research_entity else None,
                    production_type=production_type,
                    title=title[:250],
                    link=_normalize_text(item.get("link")),
                    status=normalized_status,
                    import_batch_id=(context or {}).get("import_batch_id"),
                    import_job_id=job_id or (context or {}).get("import_job_id"),
                    source_file=(context or {}).get("source_file"),
                    source_page=source_page,
                    source_section=item_source_section,
                    raw_value=raw_title or title,
                    normalized_value=normalized_title[:300],
                    confidence_score=confidence_score,
                    parser_version=(context or {}).get("parser_version") or PARSER_VERSION,
                    normalization_action="persisted" if validation_status == "validated" else "pending_review",
                    normalization_reason=review_reason or "Produccion detectada en seccion oficial de produccion cientifica.",
                    raw_title=raw_title or title,
                    normalized_title=normalized_title[:300],
                    raw_authors=raw_authors,
                    normalized_authors=normalized_authors or None,
                    raw_status=raw_status,
                    normalized_status=normalized_status,
                    raw_impact=raw_impact,
                    normalized_impact=normalized_impact,
                    validation_status=validation_status,
                    review_reason=review_reason,
                )
                self.db.add(production)
                self.db.flush()
                self._persist_production_authors(
                    production,
                    resolved_authors,
                    context or {},
                    source_page=source_page,
                    source_section=item_source_section,
                    source_field="authors",
                    row_or_block_id=row_or_block_id,
                    research_entity=research_entity,
                )
                action = "persisted" if validation_status == "validated" else "pending_review"
                reason = review_reason or "Produccion detectada en seccion oficial de produccion cientifica."
            else:
                if research_entity and not production.research_entity_id:
                    production.research_entity_id = research_entity.id
                if teacher and not production.teacher_id:
                    production.teacher_id = teacher.id
                if production.validation_status == "pending_review" and validation_status == "validated":
                    production.validation_status = "validated"
                    production.review_reason = None
                self._persist_production_authors(
                    production,
                    resolved_authors,
                    context or {},
                    source_page=source_page,
                    source_section=item_source_section,
                    source_field="authors",
                    row_or_block_id=row_or_block_id,
                    only_if_missing=True,
                    research_entity=research_entity,
                )
                action = "merged_duplicate"
                reason = "Produccion ya existia con el mismo titulo normalizado y periodo."
            self._record_normalization_audit(
                context,
                entity_type="scientific_production",
                raw_value=title,
                normalized_value=normalized_title[:300],
                action=action,
                reason=reason,
                source_section=item_source_section,
                confidence_score=confidence_score,
                normalized_record_id=production.id,
                source_page=source_page,
                metadata_json=production_metadata,
            )
            if validation_status == "pending_review":
                self._record_review_item(
                    job_id,
                    field="scientific_production",
                    raw_value=title,
                    normalized_value=normalized_title[:300],
                    reason=review_reason or "Producto cientifico requiere validacion academica.",
                    source_section=item_source_section,
                    confidence_score=confidence_score,
                    source_page=source_page,
                )
            persisted += 1
        return persisted

    @staticmethod
    def _primary_internal_author(resolved_authors: list[dict], teachers: list[Teacher]) -> Teacher | None:
        teacher_ids = {
            int(author["teacher_id"])
            for author in resolved_authors
            if author.get("author_type") == "internal" and author.get("teacher_id")
            and author.get("validation_status") == "validated"
        }
        for teacher in teachers:
            if teacher.id in teacher_ids:
                return teacher
        return None

    def _resolve_product_authors(
        self,
        authors: list[object],
        teachers: list[Teacher],
        period_id: int,
        *,
        context: dict | None = None,
        source_section: str | None = None,
        source_field: str | None = None,
        row_or_block_id: str | None = None,
        authors_source: str | None = None,
    ) -> list[dict]:
        context = context or {}
        if not _is_author_context(source_section, source_field, row_or_block_id):
            return []
        candidate_teachers = self._author_candidate_teachers(teachers, context)
        external_researchers = (
            self.db.query(ExternalResearcher)
            .filter(ExternalResearcher.period_id == period_id, ExternalResearcher.requires_review.is_(False))
            .all()
        )
        functional_roles = self._functional_author_candidates(context)
        author_mentions: list[dict] = []
        for raw_author in authors:
            author_mentions.extend(
                self._author_mentions_from_raw(
                    _normalize_text(raw_author),
                    candidate_teachers,
                    external_researchers,
                )
            )
        resolved: list[dict] = []
        for index, mention in enumerate(author_mentions, start=1):
            raw_name = str(mention.get("raw_author_name") or "")
            normalized_name = str(mention.get("normalized_author_name") or raw_name)
            preprocess_reason = str(mention.get("preprocess_reason") or "")
            possible_mixed_authors = bool(mention.get("possible_mixed_authors"))
            invalid_author_fragment = bool(mention.get("invalid_author_fragment"))
            author_tokens = _person_match_tokens(normalized_name)
            if invalid_author_fragment or len(author_tokens) < 2:
                resolved.append(
                    {
                        "author_order": index,
                        "raw_author_name": raw_name,
                        "normalized_author_name": normalized_name or None,
                        "author_type": "unresolved",
                        "teacher_id": None,
                        "external_researcher_id": None,
                        "confidence_score": 0.2,
                        "validation_status": "discarded_invalid",
                        "source_field": source_field,
                        "row_or_block_id": row_or_block_id,
                        "authors_source": authors_source,
                        "reason": preprocess_reason
                        or "Autor descartado para relacion automatica porque no parece un nombre completo confiable.",
                    }
                )
                continue

            teacher, teacher_confidence = self._match_teacher_by_tokens(candidate_teachers, author_tokens, normalized_name)
            if teacher:
                resolved.append(
                    {
                        "author_order": index,
                        "raw_author_name": raw_name,
                        "normalized_author_name": teacher.full_name,
                        "author_type": "internal",
                        "teacher_id": teacher.id,
                        "external_researcher_id": None,
                        "confidence_score": teacher_confidence,
                        "validation_status": "validated" if teacher_confidence >= 0.9 else "pending_merge_review",
                        "source_field": source_field,
                        "row_or_block_id": row_or_block_id,
                        "authors_source": authors_source,
                        "reason": _join_reasons(
                            "Autor coincide con docente interno detectado en el mismo informe.",
                            preprocess_reason,
                        ),
                    }
                )
                continue

            external, external_confidence = self._match_external_by_tokens(external_researchers, author_tokens, normalized_name)
            if external:
                resolved.append(
                    {
                        "author_order": index,
                        "raw_author_name": raw_name,
                        "normalized_author_name": external.full_name,
                        "author_type": "external",
                        "teacher_id": None,
                        "external_researcher_id": external.id,
                        "confidence_score": external_confidence,
                        "validation_status": "validated" if external_confidence >= 0.9 else "pending_merge_review",
                        "source_field": source_field,
                        "row_or_block_id": row_or_block_id,
                        "authors_source": authors_source,
                        "reason": _join_reasons(
                            "Autor coincide con investigador externo detectado para el periodo.",
                            preprocess_reason,
                        ),
                    }
                )
                continue

            functional_role, functional_confidence = self._match_functional_role_by_tokens(functional_roles, author_tokens, normalized_name)
            if functional_role:
                resolved.append(
                    {
                        "author_order": index,
                        "raw_author_name": raw_name,
                        "normalized_author_name": functional_role.normalized_name,
                        "author_type": "internal",
                        "teacher_id": functional_role.teacher_id,
                        "external_researcher_id": None,
                        "confidence_score": functional_confidence,
                        "validation_status": "validated" if functional_confidence >= 0.9 else "pending_merge_review",
                        "source_field": source_field,
                        "row_or_block_id": row_or_block_id,
                        "authors_source": authors_source,
                        "functional_role_id": functional_role.id,
                        "functional_role_type": functional_role.role_type,
                        "research_entity_id": functional_role.research_entity_id,
                        "research_project_id": functional_role.research_project_id,
                        "reason": _join_reasons(
                            "Autor coincide con persona ya detectada como director/coordinador/tutor/responsable del informe.",
                            preprocess_reason,
                        ),
                    }
                )
                continue

            pending_reason = (
                "Autor no resuelto; posible mezcla de autores en el mismo campo."
                if possible_mixed_authors
                else "Autor parece nombre valido, pero no coincide con docente interno ni investigador externo detectado."
            )
            resolved.append(
                {
                    "author_order": index,
                    "raw_author_name": raw_name,
                    "normalized_author_name": normalized_name or None,
                    "author_type": "unresolved",
                    "teacher_id": None,
                    "external_researcher_id": None,
                    "confidence_score": 0.4 if possible_mixed_authors else 0.45,
                    "validation_status": "pending_author_resolution",
                    "source_field": source_field,
                    "row_or_block_id": row_or_block_id,
                    "authors_source": authors_source,
                    "reason": _join_reasons(pending_reason, preprocess_reason),
                }
            )
        return resolved

    def _author_mentions_from_raw(
        self,
        raw_name: str,
        teachers: list[Teacher],
        external_researchers: list[ExternalResearcher],
    ) -> list[dict]:
        cleaned, ocr_reason = _clean_ocr_person_text(raw_name)
        if not cleaned:
            return []
        separator_parts = [
            part.strip()
            for part in re.split(r"\s*(?:;|\||,)\s*", cleaned)
            if part.strip()
        ]
        if len(separator_parts) > 1:
            return [
                {
                    "raw_author_name": part,
                    "normalized_author_name": _clean_person_name(part) or _title_case_name(_normalize_key(part)),
                    "preprocess_reason": _join_reasons("Autores separados por delimitador explicito.", ocr_reason),
                    "possible_mixed_authors": False,
                    "invalid_author_fragment": _invalid_author_fragment(part),
                }
                for part in separator_parts
            ]

        known_matches = self._known_author_mentions(cleaned, teachers, external_researchers)
        if len(known_matches) >= 2:
            return [
                {
                    "raw_author_name": raw_name,
                    "normalized_author_name": match["name"],
                    "preprocess_reason": _join_reasons(
                        "Campo de autores dividido por coincidencias fuertes con personas ya detectadas.",
                        ocr_reason,
                    ),
                    "possible_mixed_authors": False,
                    "invalid_author_fragment": False,
                }
                for match in known_matches
            ]

        tokens = _person_match_tokens(cleaned)
        possible_mix = len(tokens) >= 4
        return [
            {
                "raw_author_name": raw_name,
                "normalized_author_name": _clean_person_name(cleaned) or _title_case_name(_normalize_key(cleaned)),
                "preprocess_reason": ocr_reason,
                "possible_mixed_authors": possible_mix,
                "invalid_author_fragment": _invalid_author_fragment(cleaned),
            }
        ]

    @staticmethod
    def _known_author_mentions(
        raw_name: str,
        teachers: list[Teacher],
        external_researchers: list[ExternalResearcher],
    ) -> list[dict]:
        raw_tokens = _person_match_tokens(raw_name)
        if len(raw_tokens) < 6:
            return []
        matches: list[dict] = []
        for teacher in teachers:
            candidate_tokens = _person_match_tokens(teacher.full_name)
            if len(candidate_tokens) >= 2 and _person_token_match_score(candidate_tokens, raw_tokens) >= 0.98:
                matches.append({"name": teacher.full_name, "tokens": candidate_tokens})
        for external in external_researchers:
            candidate_tokens = _person_match_tokens(external.full_name)
            if len(candidate_tokens) >= 2 and _person_token_match_score(candidate_tokens, raw_tokens) >= 0.98:
                matches.append({"name": external.full_name, "tokens": candidate_tokens})
        selected: list[dict] = []
        used_tokens: set[str] = set()
        for match in sorted(matches, key=lambda item: len(item["tokens"]), reverse=True):
            overlap = used_tokens.intersection(match["tokens"])
            if overlap:
                continue
            selected.append(match)
            used_tokens.update(match["tokens"])
        return selected[:6]

    def _author_candidate_teachers(self, teachers: list[Teacher], context: dict) -> list[Teacher]:
        candidates: dict[int, Teacher] = {teacher.id: teacher for teacher in teachers}
        batch_id = context.get("import_batch_id")
        job_id = context.get("import_job_id")
        query = self.db.query(Teacher)
        if batch_id:
            for teacher in query.filter(Teacher.import_batch_id == batch_id).all():
                candidates.setdefault(teacher.id, teacher)
        if job_id:
            for teacher in query.filter(Teacher.import_job_id == job_id).all():
                candidates.setdefault(teacher.id, teacher)
        return list(candidates.values())

    def _functional_author_candidates(self, context: dict) -> list[PersonRole]:
        batch_id = context.get("import_batch_id")
        job_id = context.get("import_job_id")
        if not batch_id and not job_id:
            return []
        query = self.db.query(PersonRole).filter(
            PersonRole.role_type.in_(["director", "coordinador", "tutor_semillero", "responsable_informe"]),
            PersonRole.normalized_name.isnot(None),
        )
        if job_id:
            query = query.filter(PersonRole.import_job_id == job_id)
        elif batch_id:
            query = query.filter(PersonRole.import_batch_id == batch_id)
        return query.order_by(PersonRole.confidence_score.desc().nullslast(), PersonRole.id.asc()).all()

    @staticmethod
    def _match_teacher_by_tokens(
        teachers: list[Teacher],
        author_tokens: set[str],
        author_name: str | None = None,
    ) -> tuple[Teacher | None, float]:
        best_teacher: Teacher | None = None
        best_score = 0.0
        for teacher in teachers:
            teacher_tokens = _person_match_tokens(teacher.full_name)
            if author_tokens != teacher_tokens and not (
                author_name and heuristic_name_merge_allowed(author_name, teacher.full_name)
            ):
                continue
            score = _person_token_match_score(author_tokens, teacher_tokens)
            if score > best_score:
                best_score = score
                best_teacher = teacher
        if best_teacher and best_score >= 0.75:
            return best_teacher, round(best_score, 2)
        return None, 0.0

    @staticmethod
    def _match_external_by_tokens(
        external_researchers: list[ExternalResearcher],
        author_tokens: set[str],
        author_name: str | None = None,
    ) -> tuple[ExternalResearcher | None, float]:
        best_external: ExternalResearcher | None = None
        best_score = 0.0
        for external in external_researchers:
            external_tokens = _person_match_tokens(external.full_name)
            if author_tokens != external_tokens and not (
                author_name and heuristic_name_merge_allowed(author_name, external.full_name)
            ):
                continue
            score = _person_token_match_score(author_tokens, external_tokens)
            if score > best_score:
                best_score = score
                best_external = external
        if best_external and best_score >= 0.75:
            return best_external, round(best_score, 2)
        return None, 0.0

    @staticmethod
    def _match_functional_role_by_tokens(
        roles: list[PersonRole],
        author_tokens: set[str],
        author_name: str | None = None,
    ) -> tuple[PersonRole | None, float]:
        best_role: PersonRole | None = None
        best_score = 0.0
        for role in roles:
            role_tokens = _person_match_tokens(role.normalized_name)
            if author_tokens != role_tokens and not (
                author_name and heuristic_name_merge_allowed(author_name, role.normalized_name)
            ):
                continue
            score = _person_token_match_score(author_tokens, role_tokens)
            if score > best_score:
                best_score = score
                best_role = role
        if best_role and best_score >= 0.75:
            return best_role, round(best_score, 2)
        return None, 0.0

    def _persist_production_authors(
        self,
        production: ScientificProduction,
        resolved_authors: list[dict],
        context: dict,
        *,
        source_page: int | None,
        source_section: str,
        source_field: str | None,
        row_or_block_id: str | None,
        research_entity: ResearchEntity | None = None,
        only_if_missing: bool = False,
    ) -> None:
        persist_author_rows = True
        if only_if_missing:
            pending = any(
                isinstance(item, ScientificProductionAuthor) and item.production_id == production.id
                for item in self.db.new
            )
            if pending:
                persist_author_rows = False
            existing = (
                self.db.query(ScientificProductionAuthor.id)
                .filter(ScientificProductionAuthor.production_id == production.id)
                .first()
            )
            if existing:
                persist_author_rows = False
        else:
            self.db.query(ScientificProductionAuthor).filter(
                ScientificProductionAuthor.production_id == production.id
            ).delete(synchronize_session=False)

        for author in resolved_authors:
            if persist_author_rows:
                self.db.add(
                    ScientificProductionAuthor(
                        production_id=production.id,
                        author_order=int(author.get("author_order") or 0),
                        raw_author_name=author.get("raw_author_name"),
                        normalized_author_name=author.get("normalized_author_name"),
                        teacher_id=author.get("teacher_id"),
                        external_researcher_id=author.get("external_researcher_id"),
                        research_entity_id=author.get("research_entity_id") or (research_entity.id if research_entity else None),
                        author_type=str(author.get("author_type") or "unresolved"),
                        confidence_score=author.get("confidence_score"),
                        reason=author.get("reason"),
                        source_file=context.get("source_file"),
                        source_page=source_page,
                        source_section=source_section,
                        source_field=source_field,
                        row_or_block_id=row_or_block_id,
                        validation_status=str(author.get("validation_status") or "pending_review"),
                        metadata_json={
                            **author,
                            "row_or_block_id": row_or_block_id,
                            "source_field": source_field,
                        },
                        import_batch_id=context.get("import_batch_id"),
                        import_job_id=context.get("import_job_id"),
                        parser_version=str(context.get("parser_version") or PARSER_VERSION),
                    )
                )
            author_type = str(author.get("author_type") or "unresolved")
            if str(author.get("validation_status") or "") == "discarded_invalid":
                continue
            person_type = {
                "internal": "teacher",
                "external": "external_researcher",
                "unresolved": "unresolved",
            }.get(author_type, "unresolved")
            self._record_person_role(
                context,
                role_type="autor_producto",
                person_type=person_type,
                raw_name=author.get("raw_author_name"),
                normalized_name=author.get("normalized_author_name"),
                reason=str(author.get("reason") or "Autor detectado en fila/bloque de producto cientifico."),
                source_section=source_section,
                confidence_score=author.get("confidence_score"),
                validation_status=str(author.get("validation_status") or "pending_review"),
                period_id=production.period_id,
                teacher_id=author.get("teacher_id"),
                external_researcher_id=author.get("external_researcher_id"),
                scientific_production_id=production.id,
                research_entity_id=research_entity.id if research_entity else None,
                source_page=source_page,
                raw_value=author.get("raw_author_name"),
                normalized_value=author.get("normalized_author_name"),
                metadata_json={
                    **author,
                    "production_id": production.id,
                    "production_title": production.title,
                    "row_or_block_id": row_or_block_id,
                    "source_field": source_field,
                    "validation_status": author.get("validation_status"),
                    "source": "scientific_production_authors",
                },
            )
        if resolved_authors:
            self.db.flush()

    @staticmethod
    def _match_teacher_for_product(teachers: list[Teacher], authors: list[object]) -> Teacher | None:
        for author in authors:
            author_tokens = _person_match_tokens(author)
            if len(author_tokens) < 2:
                continue
            for teacher in teachers:
                teacher_tokens = _person_match_tokens(teacher.full_name)
                if _person_token_match_score(author_tokens, teacher_tokens) >= 0.75:
                    return teacher
        return None

    @staticmethod
    def _is_real_scientific_product_title(value: str | None) -> bool:
        return product_title_quality(value)[0] != "invalid"

    def _find_career(self, value: str | None) -> Career | None:
        if not value:
            return self.db.query(Career).order_by(Career.id.asc()).first()
        target = _normalize_key(value).replace("LICENCIATURA EN ", "").replace("LICENCIATURA ", "")
        for career in self.db.query(Career).all():
            candidate = _normalize_key(career.name).replace("LICENCIATURA EN ", "").replace("LICENCIATURA ", "")
            if target == candidate or target in candidate or candidate in target:
                return career
        return None

    @staticmethod
    def _synthetic_teacher_email(name: str, career_id: int) -> str:
        digest = hashlib.sha1(f"{name}|{career_id}".encode("utf-8")).hexdigest()[:12]
        return f"imported-{digest}@local.import"

    @staticmethod
    def _production_type_from_payload(value: object) -> ProductionType:
        key = _normalize_key(value)
        if "BOOK_CHAPTER" in key or "CAPITULO" in key:
            return ProductionType.BOOK_CHAPTER
        if "BOOK" in key or "LIBRO" in key:
            return ProductionType.BOOK
        if "PRESENTATION" in key or "PONENCIA" in key or "INTERCAMBIO" in key:
            return ProductionType.PRESENTATION
        return ProductionType.ARTICLE

    @staticmethod
    def _production_status_from_payload(value: object) -> str:
        key = _normalize_key(value)
        if any(token in key for token in ("REVISION", "ENVIADO", "SUBMITTED")):
            return "en_revision"
        return "published"

    @staticmethod
    def _progress_percent(value: object) -> int | None:
        if value is None:
            return None
        match = re.search(r"\d+", str(value))
        if not match:
            return None
        return max(0, min(100, int(match.group(0))))

    @staticmethod
    def _detect_career_name(normalized_text: str, careers: list[Career]) -> str | None:
        best_match: Career | None = None
        best_length = 0
        for career in careers:
            career_key = _normalize_key(career.name)
            short_key = career_key.replace("LICENCIATURA EN ", "")
            if career_key in normalized_text or short_key in normalized_text:
                score = len(career_key)
                if score > best_length:
                    best_match = career
                    best_length = score
        return best_match.name if best_match else None

    @staticmethod
    def _extract_director_name(text: str) -> str | None:
        patterns = [
            r"DIRECTOR DEL PROYECTO\*?\s*\n+([^\n]+)",
            r"DIRECTOR\(A\) DEL PROYECTO\*?\s*\n+([^\n]+)",
        ]
        for pattern in patterns:
            match = re.search(pattern, text, flags=re.IGNORECASE)
            if match:
                candidate = _sanitize_db_text(match.group(1))
                if candidate and "CORREO" not in _normalize_key(candidate) and _looks_like_person_name(candidate):
                    return _title_case_name(candidate)[:180]
        return None

    @staticmethod
    def _extract_participants(text: str) -> list[str]:
        marker = re.search(r"INVESTIGADORES QUE PARTICIPAN EN EL PROYECTO\*?", text, flags=re.IGNORECASE)
        if not marker:
            return []

        segment = text[marker.end() : marker.end() + 1800]
        stop = re.search(
            r"(INVESTIGADORES QUE INGRESAN|INVESTIGADORES QUE SE DESVINCULAN|RESUMEN DEL PROYECTO|SECCI[OÓ]N 2)",
            segment,
            flags=re.IGNORECASE,
        )
        if stop:
            segment = segment[: stop.start()]

        ignored = {
            "NOMBRE COMPLETO",
            "FACULTAD",
            "CARRERA",
            "FECHA DE INGRESO",
            "CIENCIAS",
            "ADMINISTRATIVAS",
            "CIENCIAS ADMINISTRATIVAS",
        }
        participants: list[str] = []
        for line in segment.splitlines():
            candidate = _sanitize_db_text(line)
            if not candidate:
                continue
            key = _normalize_key(candidate)
            if key in ignored or len(key) < 8 or re.search(r"\d", key):
                continue
            if any(word in key for word in ("LICENCIATURA", "COMERCIO", "GESTION", "CONTABILIDAD", "ADMINISTRACION")):
                continue
            if not _looks_like_person_name(candidate):
                continue
            if key not in {_normalize_key(item) for item in participants}:
                participants.append(_title_case_name(candidate)[:120])
        return participants[:12]

    @staticmethod
    def _detect_year_label(*values: str | None) -> str:
        combined = " ".join(_normalize_text(value) for value in values)
        if re.search(r"2025\s*[-/]\s*2026|CI\s*25\s*26|C[I1]\s*2526|2526", combined, flags=re.IGNORECASE):
            return "2025-2026"
        if re.search(r"2024\s*[-/]\s*2025|CI\s*24\s*25|2425", combined, flags=re.IGNORECASE):
            return "2024-2025"
        if "2025" in combined:
            return "2025-2026"
        return "2025-2026"

    @staticmethod
    def _detect_cycle(normalized_text: str) -> int:
        match = re.search(r"CICLO\s+(I{1,2}|1|2|\|)", normalized_text)
        if not match:
            return 1
        value = "I" if match.group(1) == "|" else match.group(1)
        return _roman_cycle_to_int(value)

    @staticmethod
    def _detect_articles(normalized_text: str) -> int:
        explicit = _extract_number_before(normalized_text, ["ARTICULOS", "ARTICULO"])
        if explicit:
            return min(explicit, 5)
        return 0

    @staticmethod
    def _detect_books(normalized_text: str) -> int:
        return min(_extract_number_before(normalized_text, ["LIBROS", "LIBRO"]), 3)

    @staticmethod
    def _detect_book_chapters(normalized_text: str) -> int:
        explicit = _extract_number_before(normalized_text, ["CAPITULOS", "CAPITULO"])
        if explicit:
            return min(explicit, 3)
        if "CAP DE LIBROS" in normalized_text or "CAPITULO DE LIBRO" in normalized_text:
            return 1
        return 0

    @staticmethod
    def _detect_presentations(normalized_text: str) -> int:
        explicit = _extract_number_before(normalized_text, ["PONENCIAS", "PONENCIA", "CONGRESOS", "CONGRESO"])
        if explicit:
            return min(explicit, 3)
        return 0

    def review_ocr_trace(
        self,
        trace_id: int,
        payload: OcrTraceReviewUpdate,
        reviewed_by: str | None,
    ) -> ImportedOcrTrace:
        trace = self.db.query(ImportedOcrTrace).filter(ImportedOcrTrace.id == trace_id).first()
        if not trace:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="No se encontro la traza OCR solicitada.",
            )

        allowed_statuses = {"PENDIENTE_REVISION", "VALIDADO", "VALIDADO_AUTOMATICO", "RECHAZADO"}
        if payload.review_status not in allowed_statuses:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail="El estado de revision no es valido.",
            )

        trace.review_status = payload.review_status
        trace.review_notes = _normalize_text(payload.review_notes) or None
        trace.reviewed_by = reviewed_by
        trace.reviewed_at = datetime.utcnow()
        self.db.commit()
        self.db.refresh(trace)
        return trace

    def _load_rows(self, file: UploadFile) -> dict[str, list[dict[str, object]] | list[str]]:
        if not file.filename or not file.filename.lower().endswith(".xlsx"):
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail="Solo se permiten archivos Excel .xlsx.",
            )

        content = file.file.read()
        workbook = load_workbook(BytesIO(content), data_only=True)
        worksheet = workbook[workbook.sheetnames[0]]

        raw_rows = list(worksheet.iter_rows(values_only=True))
        if not raw_rows:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail="El archivo Excel no contiene filas.",
            )

        headers = [_normalize_header(value) for value in raw_rows[0]]
        records: list[dict[str, object]] = []
        for raw in raw_rows[1:]:
            if not any(item not in (None, "") for item in raw):
                continue
            records.append({headers[index]: value for index, value in enumerate(raw) if index < len(headers)})

        return {"headers": headers, "records": records}

    @staticmethod
    def _extract_pdf_text(content: bytes) -> str:
        try:
            document = fitz.open(stream=content, filetype="pdf")
        except Exception as exc:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail="El archivo PDF no pudo abrirse correctamente.",
            ) from exc

        pages: list[str] = []
        for page in document:
            pages.append(page.get_text("text"))

        return "\n".join(part.strip() for part in pages if part and part.strip())

    @staticmethod
    def _extract_text_from_pdf(content: bytes) -> PdfTextExtraction:
        return extract_pdf_text(content)

    @staticmethod
    def _extract_pdf_text_with_tesseract(content: bytes) -> tuple[str, str | None]:
        try:
            document = fitz.open(stream=content, filetype="pdf")
        except Exception as exc:
            return "", f"El archivo PDF no pudo abrirse correctamente para OCR: {exc}"

        pages: list[str] = []
        try:
            for page in document:
                pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                image = Image.open(BytesIO(pixmap.tobytes("png")))
                text = pytesseract.image_to_string(image, lang="spa+eng")
                if text and text.strip():
                    pages.append(text.strip())
        except pytesseract.TesseractNotFoundError:
            return "", "Tesseract OCR no esta instalado o no esta disponible en el contenedor."
        except Exception as exc:
            return "", f"El OCR no pudo procesar el PDF: {exc}"

        return "\n".join(pages).strip(), None

    @staticmethod
    def _is_pdf_content(content: bytes) -> bool:
        return content.lstrip().startswith(b"%PDF-")

    @staticmethod
    def _resolve_pdf_filename(filename: str | None, source_path: str | None) -> str:
        normalized_filename = _normalize_text(filename)
        if normalized_filename.lower().endswith(".pdf"):
            return normalized_filename

        normalized_source = _normalize_text(source_path)
        if normalized_source:
            candidate = PurePosixPath(normalized_source).name
            if candidate.lower().endswith(".pdf"):
                return candidate

        if normalized_filename:
            stem = PurePosixPath(normalized_filename).stem or "progress-report"
            return f"{stem}.pdf"

        return "progress-report.pdf"

    @staticmethod
    def _validate_columns(headers: list[str], required: list[str]) -> None:
        missing = [column for column in required if column not in headers]
        if missing:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail=f"Faltan columnas requeridas: {', '.join(missing)}",
            )

    def _create_job(self, source_type: str, filename: str, imported_by: str | None) -> ImportJob:
        job = ImportJob(
            source_type=source_type,
            filename=filename,
            status="PROCESSING",
            imported_by=imported_by,
        )
        self.db.add(job)
        self.db.flush()
        return job

    def _finalize_job(
        self,
        job: ImportJob,
        imported_rows: int,
        summary: str,
        status_value: str = "SUCCESS",
    ) -> ImportResult:
        job.status = status_value
        job.summary = summary
        job.current_step = "finished"
        job.finished_at = _utcnow()
        job.duration_ms = _elapsed_ms(job.started_at, job.finished_at)
        job.processed_at = _utcnow()
        promote_document_version(self.db, job)
        self.db.commit()
        self.db.refresh(job)
        return ImportResult(
            job_id=job.id,
            batch_id=job.batch_id,
            source_type=job.source_type,
            filename=job.filename,
            status=job.status,
            imported_rows=imported_rows,
            summary=summary,
        )
